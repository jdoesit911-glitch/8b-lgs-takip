import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os
from core.analyzer import analyze_class_performance, calculate_pareto, diagnose_student, generate_support_groups

def set_cell_shading(cell, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    cell._tc.get_or_add_tcPr().append(tcMar)

def add_callout(doc, text, alert_type="NOTE"):
    colors = {
        "NOTE": ("1F497D", "F2F5F8"), "TIP": ("0E6655", "E8F8F5"),
        "IMPORTANT": ("7D6608", "FEF9E7"), "WARNING": ("900C3F", "FDEDEC")
    }
    border_col, bg_col = colors.get(alert_type, ("1F497D", "F2F5F8"))
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_shading(cell, bg_col)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    borders_xml = f'''<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="single" w:sz="24" w:space="0" w:color="{border_col}"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>'''
    cell._tc.get_or_add_tcPr().append(parse_xml(borders_xml))
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(2)
    r_tag = p.add_run(f"[{alert_type}] "); r_tag.bold = True; r_tag.font.color.rgb = RGBColor.from_string(border_col); r_tag.font.size = Pt(9.5)
    r_text = p.add_run(text); r_text.font.size = Pt(9.5); r_text.font.color.rgb = RGBColor(50, 50, 50)
    p_after = doc.add_paragraph(); p_after.paragraph_format.space_before = Pt(2); p_after.paragraph_format.space_after = Pt(2)

def generate_word_report(exam_name, students, kazanim_data=None, comparisons=None, output_path="report.docx"):
    doc = docx.Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.8); sec.bottom_margin = Inches(0.8); sec.left_margin = Inches(0.8); sec.right_margin = Inches(0.8)
    
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"; normal.font.size = Pt(10); normal.font.color.rgb = RGBColor(40, 40, 40)
    normal.paragraph_format.space_after = Pt(4); normal.paragraph_format.line_spacing = 1.15

    stats = analyze_class_performance(students)
    pareto = calculate_pareto(students)
    groups = generate_support_groups(students, kazanim_data)

    def add_title(text):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text); r.bold = True; r.font.size = Pt(18); r.font.color.rgb = RGBColor(31, 73, 125)

    def add_subtitle(text):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(14)
        r = p.add_run(text); r.italic = True; r.font.size = Pt(10.5); r.font.color.rgb = RGBColor(100, 100, 100)

    def add_h1(text):
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(3)
        r = p.add_run(text); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = RGBColor(31, 73, 125)

    def add_h2(text):
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = RGBColor(0, 128, 128)

    def add_bullet(bold_txt, regular_txt):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(0)
        r1 = p.add_run(bold_txt); r1.bold = True; r1.font.size = Pt(9.5)
        r2 = p.add_run(regular_txt); r2.font.size = Pt(9.5)

    add_title("8/B SINIFI VERİ ODAKLI AKADEMİK BAŞARI ARTIRMA VE LGS STRATEJİ SİSTEMİ")
    add_subtitle(f"Deneme Sınavı: {exam_name} | Analiz Motoru: LGS Analiz Otomasyonu v2.0\nÖrneklem: {len(students)} Öğrenci")

    add_callout(doc, f"Bu rapor {exam_name} sınav sonuçları üzerinden otomatik üretilmiştir. 38 kazanım düzeyi, resmi okul ve Türkiye geneli sıralamaları tam entegre edilmiştir.", "NOTE")

    # 1. Executive Summary
    add_h1("1. YÖNETİCİ ÖZETİ VE AKADEMİK GÖRÜNÜM")
    doc.add_paragraph(
        f"8/B sınıfı için uygulanan {exam_name} deneme sınavında sınıf mevcudu {stats.get('count', 0)} öğrenci olup, "
        f"toplam net ortalaması {stats.get('mean_net', 0.0):.2f}/90, medyan neti {stats.get('median_net', 0.0):.2f}, "
        f"puan ortalaması ise {stats.get('mean_score', 0.0):.2f} olarak hesaplanmıştır. En yüksek net {stats.get('max_net', 0.0):.2f}, "
        f"en düşük net ise {stats.get('min_net', 0.0):.2f}'dir."
    )

    # 2. Pareto Summary Table
    add_h1("2. PARETO NET KAYBI VE BRANŞ DAĞILIMI")
    tbl_p = doc.add_table(rows=len(pareto)+1, cols=7)
    tbl_p.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdrs = ["Sıra", "Ders", "Kayıp Net", "Toplam Y", "Toplam B", "Kayıp Payı %", "Küm. %"]
    for c, h in enumerate(hdrs):
        cell = tbl_p.cell(0, c)
        cell.paragraphs[0].text = h; cell.paragraphs[0].runs[0].bold = True; cell.paragraphs[0].runs[0].font.size = Pt(8.5)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255); cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_shading(cell, "595959")
    for r_idx, r in enumerate(pareto, 1):
        vals = [str(r['rank']), r['subject'], f"{r['lost_net']:.2f}", str(r['total_y']), str(r['total_b']), f"%{r['share_pct']:.1f}", f"%{r['cum_pct']:.1f}"]
        for c_idx, v in enumerate(vals):
            cell = tbl_p.cell(r_idx, c_idx)
            cell.paragraphs[0].text = v; cell.paragraphs[0].runs[0].font.size = Pt(8.5)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx != 1 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_margins(cell, 50, 50, 80, 80)
            if r_idx in [1, 2]:
                set_cell_shading(cell, "FDEDEC")

    # 3. Dynamic Support Groups
    add_h1("3. İHTİYAÇ ODAKLI DİNAMİK DESTEK GRUPLARI")
    for gname, members in groups.items():
        add_h2(f"{gname} ({len(members)} Öğrenci)")
        doc.add_paragraph("Öğrenciler: " + (", ".join(members) if members else "Bu grupta öğrenci bulunmuyor."))

    # 4. Individual Student Profiles
    add_h1("4. BİREYSEL ÖĞRENCİ AKADEMİK GELİŞİM PROFİLLERİ")
    sorted_students = sorted(students, key=lambda x: x.get('cls_rank', 999))
    for s in sorted_students:
        name = s['name']
        rank = s.get('cls_rank', '-')
        num = s.get('num', '-')
        t_net = s.get('top_net', 0.0)
        puan = s.get('puan', 0.0)
        okul_r = s.get('okul_rank', '-')
        tg_r = s.get('tg_rank', '-')
        dilim = f"%{s.get('dilim'):.2f}" if s.get('dilim') else "-"
        pen = round(s.get('top_y', 0) / 3.0, 2)
        
        diag = diagnose_student(s, kazanim_data.get(name, []) if kazanim_data else None)
        zero_str = ", ".join(diag['zero_kaz'][:4]) if diag['zero_kaz'] else "Sıfır çeken kazanım yok."
        low_str = ", ".join(diag['low_kaz'][:4]) if diag['low_kaz'] else "Kritik eksik yok."
        
        add_h2(f"{rank}. {name} (Okul No: {num} | Net: {t_net:.2f} | Puan: {puan:.2f})")
        add_bullet("Resmi Sıralama: ", f"Sınıf: {rank} | Okul: {okul_r} | Türkiye Geneli: {tg_r} | Dilim: {dilim}")
        add_bullet("Ders Dağılımı: ", f"T:{s.get('tur_net',0)} İ:{s.get('ink_net',0)} D:{s.get('din_net',0)} E:{s.get('ing_net',0)} M:{s.get('mat_net',0)} F:{s.get('fen_net',0)} | D:{s.get('top_d',0)} Y:{s.get('top_y',0)} B:{s.get('top_b',0)} | Ceza Kaybı: -{pen} Net")
        add_bullet("Sıfır Çekilen Kazanımlar: ", zero_str)
        add_bullet("Kritik Açıklar (<%50): ", low_str)
        add_bullet("Hedefler ve Strateji: ", f"14. Gün: {diag['target_14d']} Net | 28. Gün: {diag['target_28d']} Net. {diag['strategy']}")
        add_bullet("Haftalık Görev Kotası: ", f"{diag['weekly_quota']} soru + Yanlış Defteri Takibi.")

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc.save(output_path)
    return output_path
