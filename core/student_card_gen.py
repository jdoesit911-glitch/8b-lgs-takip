import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os
import zipfile
from core.analyzer import diagnose_student
from core.youtube_recommender import get_youtube_advice

def norm_text(text):
    if not text:
        return ""
    tr_map = {
        'İ': 'i', 'I': 'i', 'ı': 'i', 'i': 'i',
        'Ğ': 'g', 'ğ': 'g',
        'Ü': 'u', 'ü': 'u',
        'Ş': 's', 'ş': 's',
        'Ö': 'o', 'ö': 'o',
        'Ç': 'c', 'ç': 'c'
    }
    t = str(text)
    for k, v in tr_map.items():
        t = t.replace(k, v)
    return t.lower()

def set_cell_shd(cell, hex_col):
    cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_col}"/>'))

def set_cell_margins(cell, top=70, bottom=70, left=100, right=100):
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    cell._tc.get_or_add_tcPr().append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

mat_names = {
    "Verilen Pozitif Tam Sayıların Pozitif Tam Sayı Çarpanlarını": "Çarpanlar ve Katlar / Asal Çarpanlar",
    "İki Doğal Sayının En Büyük Ortak Bölenini": "EBOB-EKOK Problemleri",
    "Üslü İfadelerle İlgili Temel Kuralları": "Üslü İfadeler Temel Kuralları",
    "Verilen İki Doğal Sayının Aralarında Asal": "Aralarında Asal Sayılar",
    "Tam Sayıların, Tam Sayı Kuvvetlerini": "Tam Sayı Kuvvetleri (Negatif Üs)"
}

din_names = {
    "1. Kader ve Kaza": "Kader ve Kaza İnancı",
    "2. İnsanın İradesi": "İnsanın İradesi ve Kader",
    "3. Kaderle İlgili": "Kader Kavramları (Ömür, Ecel, Rızık, Tevekkül)",
    "4. Bir Peygamber": "Hz. Musa (A.S.) ve Mücadelesi",
    "5. Bir Ayet": "Ayet el-Kürsi ve Anlamı"
}

ink_names = {
    "Mavi Gözlü Çocuk": "Mustafa Kemal'in Öğrenim Hayatı",
    "Uyanan Avrupa": "Uyanan Avrupa ve Sarsılan Osmanlı",
    "Buhranlar Büyük": "Mustafa Kemal'in Fikir Hayatını Etkileyen Olaylar",
    "Adım Adım": "Mustafa Kemal'in Askerlik Hayatı"
}

ing_names = {
    "Invitation Cards": "Davet Kartları / Mektupları (Invitation)",
    "Accepting & Refusing": "Daveti Kabul & Reddetme Kalıpları",
    "Making Simple": "Basit Sorular Sorma (Simple Inquiries)",
    "Understanding Texts": "Friendship Okuduğunu Anlama",
    "Some Qualities": "Kişisel Özellikler / Sıfatlar (Qualities)",
    "Making an Excuse": "Mazeret Bildirme ve Özür Dileme"
}

def clean_kaz_name(sub, name):
    sub_l = norm_text(sub)
    name_l = norm_text(name)
    if "mat" in sub_l:
        for k, v in mat_names.items():
            if norm_text(k) in name_l: return v
    elif "din" in sub_l:
        for k, v in din_names.items():
            if norm_text(k) in name_l: return v
    elif "ink" in sub_l or "tarih" in sub_l:
        for k, v in ink_names.items():
            if norm_text(k) in name_l: return v
    elif "ing" in sub_l or "dil" in sub_l:
        for k, v in ing_names.items():
            if norm_text(k) in name_l: return v
    return name

def get_subject_prescriptions(student, kazanim_list):
    """
    Returns a dictionary of detailed, highly personalized prescriptions for all 6 LGS subjects:
    - Fen Bilimleri (2 kazanım)
    - Matematik (5 kazanım)
    - Türkçe (16 kazanım)
    - T.C. İnkılap Tarihi ve Atatürkçülük (4 kazanım)
    - Din Kültürü ve Ahlak Bilgisi (5 kazanım)
    - Yabancı Dil (İngilizce) (6 kazanım)
    """
    prescriptions = {}

    def get_sub_k(sub_kw):
        res = []
        kw = norm_text(sub_kw)
        for k in kazanim_list:
            if kw in norm_text(k.get('subject', '')):
                res.append(k)
        return res

    # 1. FEN BİLİMLERİ (2 Kazanım)
    f_net = float(student.get('fen_net', 0.0) or 0.0)
    f_d, f_y, f_b = int(student.get('fen_d',0)), int(student.get('fen_y',0)), int(student.get('fen_b',0))
    f_k = get_sub_k("fen")
    
    f_lines = [f"{clean_kaz_name('fen', k['kazanim'])}: %{k['rate']:.1f}" for k in f_k]
    f_summary = " | ".join(f_lines) if f_lines else "Kazanım başarısı kaydedildi."
    
    f_zero = [clean_kaz_name('fen', k['kazanim']) for k in f_k if k['rate'] == 0]
    f_low = [f"{clean_kaz_name('fen', k['kazanim'])} (%{k['rate']:.0f})" for k in f_k if 0 < k['rate'] < 100]
    f_mast = [clean_kaz_name('fen', k['kazanim']) for k in f_k if k['rate'] == 100]
    
    f_mevsim = next((k for k in f_k if "mevsim" in norm_text(k.get('kazanim',''))), None)
    f_iklim = next((k for k in f_k if "iklim" in norm_text(k.get('kazanim',''))), None)

    f_adv = f"Fen Bilimleri netin: {f_net:.2f} ({f_d} Doğru, {f_y} Yanlış, {f_b} Boş). "
    if f_net >= 16.0:
        f_adv += "Kavramsal kavrayışın ve deney yorumlama becerin çok yüksek. "
        if f_mevsim and f_mevsim['rate'] < 100:
            f_adv += f"Tek hata yaptığın 'Mevsimlerin Oluşumu'ndaki (%{f_mevsim['rate']:.0f}) eksen eğikliği ve Güneş ışınlarının geliş açısı grafiklerini pekiştirerek 20'de 20 full hedefini yakalayabilirsin."
        elif f_iklim and f_iklim['rate'] < 100:
            f_adv += f"Tek eksiğin olan 'İklim ve Hava Hareketleri' (%{f_iklim['rate']:.0f}) rüzgar oluşumu ve basınç farkı grafiklerine dikkat etmelisin."
        else:
            f_adv += "Her iki ünitede de (%100) tam hakimiyet sergiledin. Çok değişkenli yeni nesil sorularla formunu koru."
    elif f_net >= 9.0:
        f_adv += f"Sağlam bir altyapın var ancak {f_y} yanlış netini aşağı çekmiş. "
        if f_mevsim and f_mevsim['rate'] < 75:
            f_adv += f"Özellikle 'Mevsimlerin Oluşumu' konusundaki (%{f_mevsim['rate']:.0f}) 21 Haziran/21 Aralık gün dönümleri ve birim alana düşen enerji miktarı şemalarını 3D simülasyonla tekrar etmelisin. "
        if f_iklim and f_iklim['rate'] < 75:
            f_adv += f"'İklim ve Hava Hareketleri' konusunda (%{f_iklim['rate']:.0f}) yüksek basınçtan alçak basınca yatay hava akımı (rüzgar) ilişkisini netleştirmelisin. "
        f_adv += "Deney sorularında bağımsız değişkeni kenara yaz; haftalık 45 fen sorusu çöz."
    else:
        f_adv += f"Fen Bilimleri dersi öncelikli müdahale alanındır ({f_net:.2f} Net, {f_y} Yanlış). "
        if f_y >= 5: f_adv += "Kesin Kural: Emin olmadığın soruları rastgele tahminle işaretleme; bilmediğini boş bırakmak dahi netini hemen yükseltecektir. "
        if f_mevsim and f_mevsim['rate'] < 50: f_adv += f"'Mevsimlerin Oluşumu' (%{f_mevsim['rate']:.0f}) temel kavramlarını sıfırdan çalış. "
        if f_iklim and f_iklim['rate'] < 50: f_adv += f"'İklim ve Hava Hareketleri' (%{f_iklim['rate']:.0f}) hava olayları kavram haritasını çıkar. "
        f_adv += "Haftalık 30 temel seviye kazanım sorusu çöz ve yapamadığın her soruyu öğretmenine sor."

    f_target = f_zero[0] if f_zero else (f_low[0].split(' (')[0] if f_low else None)
    f_yt = get_youtube_advice("Fen Bilimleri", f_net, f_target)

    prescriptions["Fen Bilimleri"] = {
        "title": "Fen Bilimleri Gelişim Reçetesi",
        "icon": "[FEN BİLİMLERİ]",
        "net": f_net, "d": f_d, "y": f_y, "b": f_b,
        "summary": f_summary, "zero": f_zero, "low": f_low, "mast": f_mast, "advice": f_adv,
        "youtube": f_yt
    }

    # 2. MATEMATİK (5 Kazanım)
    m_net = float(student.get('mat_net', 0.0) or 0.0)
    m_d, m_y, m_b = int(student.get('mat_d',0)), int(student.get('mat_y',0)), int(student.get('mat_b',0))
    m_k = get_sub_k("mat")
    m_lines = [f"{clean_kaz_name('mat', k['kazanim'])}: %{k['rate']:.0f}" for k in m_k]
    m_summary = " | ".join(m_lines) if m_lines else "Kazanım başarısı kaydedildi."
    
    m_zero = [clean_kaz_name('mat', k['kazanim']) for k in m_k if k['rate'] == 0]
    m_low = [f"{clean_kaz_name('mat', k['kazanim'])} (%{k['rate']:.0f})" for k in m_k if 0 < k['rate'] < 100]
    m_mast = [clean_kaz_name('mat', k['kazanim']) for k in m_k if k['rate'] == 100]
    
    m_adv = f"Matematik netin: {m_net:.2f} ({m_d} Doğru, {m_y} Yanlış, {m_b} Boş). "
    if m_net >= 14.0:
        m_adv += "Yeni nesil problem çözme ve modelleme becerin okul zirvesindedir. "
        if m_zero or m_low:
            missed = m_zero + [x.split(' (')[0] for x in m_low]
            m_adv += f"Sınavda tek net kaybının yaşandığı '{', '.join(missed)}' konusundaki uç soru kalıplarına odaklanarak 20'de 20 full yapabilirsin. Süreli branş denemelerine devam et."
        else:
            m_adv += "5 kazanımın tamamında (%100) kusursuz başarı sağladın. Hız ve işlem kontrol rutinini sürdür."
    elif m_net >= 6.0:
        m_adv += f"Temel matematik altyapın oluşmuş ancak {m_y} yanlış ve {m_b} boş puanını sınırlamış. "
        if m_zero:
            m_adv += f"Sıfır çektiğin '{', '.join(m_zero)}' konusunu konu anlatım fasikülünden baştan tara. "
        if m_low:
            m_adv += f"Kritik açıkların olan [{', '.join(m_low)}] konularında 'Verilenleri Yaz -> Şemayı Çiz -> Matematiksel Modeli Kur' 3 adımlı protokolünü uygula. "
        m_adv += "Günde en az 20 adet yeni nesil soru çözerek işlem hızını artır."
    else:
        m_adv += f"Matematik dersi LGS başarında en kritik sıçrama alanındır ({m_net:.2f} Net, {m_y} Yanlış). "
        if m_y >= 5: m_adv += "Stratejik Kural: Bilmediğin soruyu boş bırakmak 0.33 net kazandırır; rastgele tahmin yapmayı derhal bırak! "
        if m_zero or m_low:
            targets = (m_zero + [x.split(' (')[0] for x in m_low])[:2]
            m_adv += f"Öncelikle [{', '.join(targets)}] konularındaki temel kazanım kavrama testlerini çöz. "
    m_target = m_zero[0] if m_zero else (m_low[0].split(' (')[0] if m_low else None)
    m_yt = get_youtube_advice("Matematik", m_net, m_target)

    prescriptions["Matematik"] = {
        "title": "Matematik Gelişim Reçetesi",
        "icon": "[MATEMATİK]",
        "net": m_net, "d": m_d, "y": m_y, "b": m_b,
        "summary": m_summary, "zero": m_zero, "low": m_low, "mast": m_mast, "advice": m_adv,
        "youtube": m_yt
    }

    # 3. TÜRKÇE (16 Kazanım)
    t_net = float(student.get('tur_net', 0.0) or 0.0)
    t_d, t_y, t_b = int(student.get('tur_d',0)), int(student.get('tur_y',0)), int(student.get('tur_b',0))
    t_k = get_sub_k("turkce")
    
    t_zero = [k['kazanim'] for k in t_k if k['rate'] == 0]
    t_low = [f"{k['kazanim']} (%{k['rate']:.0f})" for k in t_k if 0 < k['rate'] < 100]
    t_mast = [k['kazanim'] for k in t_k if k['rate'] == 100]
    
    if len(t_mast) == len(t_k) and len(t_k) > 0:
        t_summary = f"16 kazanımın tamamında %100 tam başarı sağlandı."
    else:
        imperfect = [f"{k['kazanim']}: %{k['rate']:.0f}" for k in t_k if k['rate'] < 100]
        t_summary = f"{len(t_mast)}/16 Tam Başarı | Gelişim Alanları: " + " | ".join(imperfect[:4])
        
    t_adv = f"Türkçe branş netin: {t_net:.2f} ({t_d} Doğru, {t_y} Yanlış, {t_b} Boş). "
    if t_net >= 15.0:
        t_adv += "Okuduğunu anlama, çıkarım ve muhakeme gücün çok yüksek. "
        if t_zero:
            t_adv += f"Sıfır çektiğin '{', '.join(t_zero)}' konusundaki soru köklerine ve çeldiricilere dikkat ederek 20'de 20 yapabilirsin."
        elif t_low:
            t_adv += f"Kısmi kayıp yaşadığın [{', '.join(t_low[:2])}] alanlarını gözden geçirerek 20'de 20 full hedefini yakala."
        else:
            t_adv += "Formunu korumak için her gün 20 süreli paragraf sorusu çözmeye devam et."
    elif t_net >= 10.0:
        t_adv += f"Ortalama üstü bir performans olmakla birlikte {t_y} yanlış başarı payını törpülemiş. "
        if t_zero:
            t_adv += f"Öncelikli olarak sıfır çektiğin [{', '.join(t_zero[:2])}] konularındaki kavram ve kuralları tekrar et. "
        if t_low:
            t_adv += f"Kısmi kayıp yaşadığın [{', '.join(t_low[:2])}] alanlarında metindeki açık kanıtın altını çizerek çöz. "
        t_adv += "Günde 20 paragraf sorusu çöz; iki şık arasında kaldığında işaretleme yapmayıp boş bırak."
    else:
        t_adv += f"Türkçe branşında {t_y} yanlış ciddi bir net kaybına yol açmış ({t_net:.2f} Net). "
        if t_zero:
            t_adv += f"Acilen sıfır çektiğin [{', '.join(t_zero[:3])}] konularını konu fasikülünden çalış. "
        t_adv += "Paragrafı kendi yorumunla değil metindeki kanıtla çöz; günde 20 paragraf + 10 dil bilgisi çöz."

    t_target = t_zero[0] if t_zero else (t_low[0].split(' (')[0] if t_low else None)
    t_yt = get_youtube_advice("Türkçe", t_net, t_target)

    prescriptions["Türkçe"] = {
        "title": "Türkçe Gelişim Reçetesi",
        "icon": "[TÜRKÇE]",
        "net": t_net, "d": t_d, "y": t_y, "b": t_b,
        "summary": t_summary, "zero": t_zero, "low": t_low, "mast": t_mast, "advice": t_adv,
        "youtube": t_yt
    }

    # 4. İNKILAP TARİHİ (4 Kazanım)
    i_net = float(student.get('ink_net', 0.0) or 0.0)
    i_d, i_y, i_b = int(student.get('ink_d',0)), int(student.get('ink_y',0)), int(student.get('ink_b',0))
    i_k = get_sub_k("ink") or get_sub_k("tarih")
    i_lines = [f"{clean_kaz_name('ink', k['kazanim'])}: %{k['rate']:.0f}" for k in i_k]
    i_summary = " | ".join(i_lines) if i_lines else "Kazanım başarısı kaydedildi."
    
    i_zero = [clean_kaz_name('ink', k['kazanim']) for k in i_k if k['rate'] == 0]
    i_low = [f"{clean_kaz_name('ink', k['kazanim'])} (%{k['rate']:.0f})" for k in i_k if 0 < k['rate'] < 100]
    i_mast = [clean_kaz_name('ink', k['kazanim']) for k in i_k if k['rate'] == 100]
    
    i_adv = f"İnkılap Tarihi netin: {i_net:.2f} ({i_d} Doğru, {i_y} Yanlış, {i_b} Boş). "
    if i_net >= 8.0:
        i_adv += "Tarihsel kronoloji ve neden-sonuç analiz kabiliyetin çok güçlü. "
        if i_zero or i_low:
            missed = i_zero + [x.split(' (')[0] for x in i_low]
            i_adv += f"Tek net kaybının yaşandığı '{', '.join(missed)}' konusundaki kavram haritasını inceleyerek 10'da 10 full hedefini sürdür."
        else:
            i_adv += "4 kazanımın tamamında %100 tam başarı sağladın. Haftalık 15 soruyla aralıklı tekrar yap."
    else:
        i_adv += f"Hedeflenen seviyenin altında kalındı ({i_net:.2f} Net). "
        if i_zero:
            i_adv += f"Sıfır çektiğin [{', '.join(i_zero)}] konusundaki kavram haritalarını çalış. "
        if i_low:
            i_adv += f"Pekiştirilmesi gereken [{', '.join(i_low)}] konularında kronoloji tablosu çıkar. "
        i_adv += "Sorularda 'Yalnızca bu bilgiye göre' köklerine dikkat et; metne sadık kal. Haftalık 25 soru çöz."

    i_target = i_zero[0] if i_zero else (i_low[0].split(' (')[0] if i_low else None)
    i_yt = get_youtube_advice("İnkılap Tarihi", i_net, i_target)

    prescriptions["İnkılap Tarihi"] = {
        "title": "T.C. İnkılap Tarihi ve Atatürkçülük Reçetesi",
        "icon": "[İNKILAP TARİHİ]",
        "net": i_net, "d": i_d, "y": i_y, "b": i_b,
        "summary": i_summary, "zero": i_zero, "low": i_low, "mast": i_mast, "advice": i_adv,
        "youtube": i_yt
    }

    # 5. DİN KÜLTÜRÜ (5 Kazanım)
    d_net = float(student.get('din_net', 0.0) or 0.0)
    d_d, d_y, d_b = int(student.get('din_d',0)), int(student.get('din_y',0)), int(student.get('din_b',0))
    d_k = get_sub_k("din")
    d_lines = [f"{clean_kaz_name('din', k['kazanim'])}: %{k['rate']:.0f}" for k in d_k]
    d_summary = " | ".join(d_lines) if d_lines else "Kazanım başarısı kaydedildi."
    
    d_zero = [clean_kaz_name('din', k['kazanim']) for k in d_k if k['rate'] == 0]
    d_low = [f"{clean_kaz_name('din', k['kazanim'])} (%{k['rate']:.0f})" for k in d_k if 0 < k['rate'] < 100]
    d_mast = [clean_kaz_name('din', k['kazanim']) for k in d_k if k['rate'] == 100]
    
    d_adv = f"Din Kültürü netin: {d_net:.2f} ({d_d} Doğru, {d_y} Yanlış, {d_b} Boş). "
    if d_net >= 8.0:
        d_adv += "Kavram ve ayet yorumlama hakimiyetin yüksek. "
        if d_zero or d_low:
            missed = d_zero + [x.split(' (')[0] for x in d_low]
            d_adv += f"Tek net kaybının yaşandığı '{', '.join(missed)}' konusundaki kavram veya ayet mealini inceleyerek 10'da 10 yapabilirsin."
        else:
            d_adv += "5 kazanımın tamamında %100 tam başarı sağladın. Formunu haftalık 15 soruyla koru."
    else:
        d_adv += f"Din Kültürü netin ({d_net:.2f}) potansiyelinin altında. "
        if d_zero:
            d_adv += f"Sıfır çektiğin [{', '.join(d_zero)}] konusunu mutlaka tekrar et. "
        if d_low:
            d_adv += f"Kavram karmaşası yaşanan [{', '.join(d_low)}] alanlarına odaklan. "
        d_adv += "Ömür, ecel, rızık, tevekkül, cüzi/külli irade kavramlarını sözlük kartlarına yazıp ezberle. Ayet-el Kürsi ana fikrine ve yasalara çalış."

    d_target = d_zero[0] if d_zero else (d_low[0].split(' (')[0] if d_low else None)
    d_yt = get_youtube_advice("Din Kültürü", d_net, d_target)

    prescriptions["Din Kültürü"] = {
        "title": "Din Kültürü ve Ahlak Bilgisi Reçetesi",
        "icon": "[DİN KÜLTÜRÜ]",
        "net": d_net, "d": d_d, "y": d_y, "b": d_b,
        "summary": d_summary, "zero": d_zero, "low": d_low, "mast": d_mast, "advice": d_adv,
        "youtube": d_yt
    }

    # 6. İNGİLİZCE (6 Kazanım)
    e_net = float(student.get('ing_net', 0.0) or 0.0)
    e_d, e_y, e_b = int(student.get('ing_d',0)), int(student.get('ing_y',0)), int(student.get('ing_b',0))
    e_k = get_sub_k("ing") or get_sub_k("dil")
    e_lines = [f"{clean_kaz_name('ing', k['kazanim'])}: %{k['rate']:.0f}" for k in e_k]
    e_summary = " | ".join(e_lines) if e_lines else "Kazanım başarısı kaydedildi."
    
    e_zero = [clean_kaz_name('ing', k['kazanim']) for k in e_k if k['rate'] == 0]
    e_low = [f"{clean_kaz_name('ing', k['kazanim'])} (%{k['rate']:.0f})" for k in e_k if 0 < k['rate'] < 100]
    e_mast = [clean_kaz_name('ing', k['kazanim']) for k in e_k if k['rate'] == 100]
    
    e_adv = f"İngilizce branş netin: {e_net:.2f} ({e_d} Doğru, {e_y} Yanlış, {e_b} Boş). "
    if e_net >= 8.0:
        e_adv += "Friendship ünitesi kelime ve diyalog hakimiyetin çok iyi. "
        if e_zero or e_low:
            missed = e_zero + [x.split(' (')[0] for x in e_low]
            e_adv += f"Tek net kaybının yaşandığı '{', '.join(missed)}' kalıbını gözden geçirerek 10'da 10 yapabilirsin."
        else:
            e_adv += "6 kazanımın tamamında %100 tam başarı sağladın. Soru köklerindeki olumsuz ifadelere (NOT, except) dikkat et."
    else:
        e_adv += f"İngilizce branşında {e_net:.2f} net ile hedefin gerisinde kalındı. "
        if e_zero:
            e_adv += f"Sıfır çektiğin [{', '.join(e_zero)}] soru kalıplarını tekrar et. "
        if e_low:
            e_adv += f"Pekiştirilmesi gereken [{', '.join(e_low)}] kalıplarını çalış. "
        e_adv += "Davet etme (Would you like...), kabul/ret ve mazeret kalıplarını ve Friendship sıfatlarını her gün 5 kelime kartıyla çalış. Rastgele işaretleme yapma."

    e_target = e_zero[0] if e_zero else (e_low[0].split(' (')[0] if e_low else None)
    e_yt = get_youtube_advice("İngilizce", e_net, e_target)

    prescriptions["İngilizce"] = {
        "title": "Yabancı Dil (İngilizce) Reçetesi",
        "icon": "[İNGİLİZCE]",
        "net": e_net, "d": e_d, "y": e_y, "b": e_b,
        "summary": e_summary, "zero": e_zero, "low": e_low, "mast": e_mast, "advice": e_adv,
        "youtube": e_yt
    }

    return prescriptions

def generate_single_student_card(student, exam_name, kazanim_list, output_file, comp_data=None):
    doc = docx.Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.4); sec.bottom_margin = Inches(0.4)
    sec.left_margin = Inches(0.5); sec.right_margin = Inches(0.5)
    
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"; normal.font.size = Pt(8.5)
    normal.paragraph_format.space_after = Pt(2); normal.paragraph_format.line_spacing = 1.1

    # Extract clean student metadata with aliases
    name = student['name'].strip()
    no = str(student.get('num') or student.get('school_no') or '---').strip()
    cls_name = student.get('class_name', '8/B')
    t_net = float(student.get('top_net', 0.0) or 0.0)
    puan = float(student.get('puan') if student.get('puan') is not None else student.get('score', 0.0) or 0.0)
    cls_r = student.get('cls_rank') if student.get('cls_rank') is not None else student.get('class_rank', '-')
    okul_r = student.get('okul_rank') if student.get('okul_rank') is not None else student.get('school_rank', '-')
    tg_r = student.get('tg_rank') or '-'
    
    dilim_raw = student.get('dilim') if student.get('dilim') is not None else student.get('tg_percentile')
    dilim_str = f"%{float(dilim_raw):.2f}" if dilim_raw else "-"

    total_cls = student.get('total_class_students', 28)
    total_sch = student.get('total_school_students', 56 if 'NARTEST' in str(exam_name).upper() else 110)

    # Header
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(2)
    p_title.paragraph_format.space_after = Pt(1)
    r_t = p_title.add_run(f"{cls_name} SINIFI BİREYSEL LGS GELİŞİM VE HEDEF KARNESİ")
    r_t.bold = True; r_t.font.size = Pt(12.5); r_t.font.color.rgb = RGBColor(31, 73, 125)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(3)
    r_s = p_sub.add_run(f"Deneme Sınavı: {exam_name} | Kurumsal Başarı ve Takip Sistemi")
    r_s.italic = True; r_s.font.size = Pt(8.5); r_s.font.color.rgb = RGBColor(100, 100, 100)

    # Student Info Table
    tbl_info = doc.add_table(rows=2, cols=4)
    tbl_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_info, color="B0C4DE", sz="4")
    labels = [
        ("Öğrenci Adı Soyadı:", name),
        ("Okul Numarası / Şube:", f"{no} / {cls_name}"),
        ("Toplam Net / Puan:", f"{t_net:.2f} Net / {puan:.2f} Puan"),
        ("Resmi Başarı Sıralaması:", f"Sınıf: {cls_r} / {total_cls} | Okul: {okul_r} / {total_sch} | TG: {tg_r} (Dilim: {dilim_str})")
    ]
    for idx, (lbl, val) in enumerate(labels):
        r_i = idx // 2; c_i = (idx % 2) * 2
        c1 = tbl_info.cell(r_i, c_i); c2 = tbl_info.cell(r_i, c_i+1)
        c1.paragraphs[0].text = lbl; c1.paragraphs[0].runs[0].bold = True
        c1.paragraphs[0].runs[0].font.size = Pt(8)
        c2.paragraphs[0].text = val
        c2.paragraphs[0].runs[0].font.size = Pt(8)
        set_cell_shd(c1, "DCE6F1")
        set_cell_margins(c1, 35, 35, 50, 50)
        set_cell_margins(c2, 35, 35, 50, 50)

    doc.add_paragraph().paragraph_format.space_after = Pt(1)

    # Subject Results Table
    tbl_res = doc.add_table(rows=8, cols=7)
    tbl_res.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_res, color="B0C4DE", sz="4")
    hdrs = ["Ders Adı", "Soru", "Doğru", "Yanlış", "Boş", "Net", "Başarı %"]
    for c, h in enumerate(hdrs):
        cell = tbl_res.cell(0, c)
        cell.paragraphs[0].text = h; cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(7.5)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_shd(cell, "1F497D"); cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(cell, 25, 25, 35, 35)

    subs = [
        ("Türkçe", 20, student.get('tur_d',0), student.get('tur_y',0), student.get('tur_b',0), float(student.get('tur_net',0.0)), (float(student.get('tur_net',0.0))/20.0)*100),
        ("İnkılap Tarihi", 10, student.get('ink_d',0), student.get('ink_y',0), student.get('ink_b',0), float(student.get('ink_net',0.0)), (float(student.get('ink_net',0.0))/10.0)*100),
        ("Din Kültürü", 10, student.get('din_d',0), student.get('din_y',0), student.get('din_b',0), float(student.get('din_net',0.0)), (float(student.get('din_net',0.0))/10.0)*100),
        ("İngilizce", 10, student.get('ing_d',0), student.get('ing_y',0), student.get('ing_b',0), float(student.get('ing_net',0.0)), (float(student.get('ing_net',0.0))/10.0)*100),
        ("Matematik", 20, student.get('mat_d',0), student.get('mat_y',0), student.get('mat_b',0), float(student.get('mat_net',0.0)), (float(student.get('mat_net',0.0))/20.0)*100),
        ("Fen Bilimleri", 20, student.get('fen_d',0), student.get('fen_y',0), student.get('fen_b',0), float(student.get('fen_net',0.0)), (float(student.get('fen_net',0.0))/20.0)*100),
        ("TOPLAM", 90, student.get('top_d',0), student.get('top_y',0), student.get('top_b',0), t_net, (t_net/90.0)*100)
    ]
    for r_idx, sdata in enumerate(subs, 1):
        for c_idx, val in enumerate(sdata):
            cell = tbl_res.cell(r_idx, c_idx)
            if c_idx == 6:
                cell.paragraphs[0].text = f"%{max(val, 0.0):.1f}"
            elif isinstance(val, float):
                cell.paragraphs[0].text = f"{val:.2f}"
            else:
                cell.paragraphs[0].text = str(val)
                
            cell.paragraphs[0].runs[0].font.size = Pt(7.5)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            set_cell_margins(cell, 20, 20, 35, 35)
            if r_idx == 7:
                cell.paragraphs[0].runs[0].bold = True; set_cell_shd(cell, "DCE6F1")

    # ─── DENEMELER ARASI GELİŞİM VE KARŞILAŞTIRMA BÖLÜMÜ ───
    has_comp_section = False
    if comp_data:
        st_comp = comp_data.get(name)
        if st_comp and st_comp.get('status') == 'normal':
            has_comp_section = True
            p_cmp_hdr = doc.add_paragraph()
            p_cmp_hdr.paragraph_format.space_before = Pt(4)
            p_cmp_hdr.paragraph_format.space_after = Pt(2)
            r_ch = p_cmp_hdr.add_run("1. DENEMELER ARASI GELİŞİM VE AKADEMİK DEĞİŞİM ANALİZİ (PARAF 1 ➔ NARTEST 1)")
            r_ch.bold = True; r_ch.font.size = Pt(9); r_ch.font.color.rgb = RGBColor(31, 73, 125)

            # Overview comparison table (5 rows x 5 cols)
            tbl_cmp = doc.add_table(rows=5, cols=5)
            tbl_cmp.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(tbl_cmp, color="B0C4DE", sz="4")
            cmp_hdrs = ["Değerlendirme Kriteri", "1. Deneme (PARAF)", "2. Deneme (NARTEST)", "Sınavlar Arası Fark", "Gelişim Durumu"]
            for c_i, h in enumerate(cmp_hdrs):
                c = tbl_cmp.cell(0, c_i)
                c.paragraphs[0].text = h; c.paragraphs[0].runs[0].bold = True
                c.paragraphs[0].runs[0].font.size = Pt(7.5)
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if c_i == 0 else WD_ALIGN_PARAGRAPH.CENTER
                set_cell_shd(c, "1F497D"); c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
                set_cell_margins(c, 25, 25, 35, 35)

            r1 = st_comp['r1']
            r2 = st_comp['r2']
            d_net = st_comp['delta_net']
            d_puan = st_comp['delta_puan']
            d_cls = st_comp['delta_cls_rank']
            d_sch = st_comp['delta_sch_rank']

            diff_net_str = f"+{d_net:.2f} Net" if d_net >= 0 else f"{d_net:.2f} Net"
            diff_puan_str = f"+{d_puan:.2f} Puan" if d_puan >= 0 else f"{d_puan:.2f} Puan"
            diff_cls_str = f"+{d_cls} Sıra Yükseldi" if d_cls > 0 else (f"{d_cls} Sıra Geriledi" if d_cls < 0 else "Sırası Değişmedi")
            diff_sch_str = f"+{d_sch} Sıra Yükseldi" if d_sch > 0 else (f"{d_sch} Sıra Geriledi" if d_sch < 0 else "Sırası Değişmedi")

            cmp_rows = [
                ("Toplam Net", f"{r1['top_net']:.2f} Net", f"{r2['top_net']:.2f} Net", diff_net_str, f"{st_comp['trend_icon']} {st_comp['trend']}"),
                ("LGS Puanı", f"{r1['puan']:.2f} Puan", f"{r2['puan']:.2f} Puan", diff_puan_str, "Belirgin Puan Artışı" if d_puan > 15 else ("Puan Artışı" if d_puan > 0 else "Puan Kaybı")),
                ("Sınıf Derecesi", f"{r1['class_rank']}. Sıra / 27", f"{r2['class_rank']}. Sıra / 28", diff_cls_str, "Sınıf İçi Başarı"),
                ("Okul Derecesi", f"{r1['school_rank']}. Sıra / 110", f"{r2['school_rank']}. Sıra / 56", diff_sch_str, "Okul İçi Başarı")
            ]
            for r_i, rdata in enumerate(cmp_rows, 1):
                for c_i, val in enumerate(rdata):
                    cell = tbl_cmp.cell(r_i, c_i)
                    cell.paragraphs[0].text = val
                    cell.paragraphs[0].runs[0].font.size = Pt(7.5)
                    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if c_i == 0 else WD_ALIGN_PARAGRAPH.CENTER
                    set_cell_margins(cell, 20, 20, 35, 35)
                    if c_i == 3:
                        cell.paragraphs[0].runs[0].bold = True
                        if "+" in val: set_cell_shd(cell, "E8F8F5")
                        elif "-" in val: set_cell_shd(cell, "FDEDEC")
                    elif r_i == 1:
                        set_cell_shd(cell, "F0F4F8")

            # Horizontal 6-branch net changes table
            p_sub_hdr = doc.add_paragraph()
            p_sub_hdr.paragraph_format.space_before = Pt(3)
            p_sub_hdr.paragraph_format.space_after = Pt(1)
            r_sbh = p_sub_hdr.add_run("Branş Bazlı Net Değişim Tablosu:")
            r_sbh.bold = True; r_sbh.font.size = Pt(8); r_sbh.font.color.rgb = RGBColor(0, 102, 102)

            tbl_sub_cmp = doc.add_table(rows=3, cols=7)
            tbl_sub_cmp.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(tbl_sub_cmp, color="B0C4DE", sz="4")
            sub_cmp_hdrs = ["Deneme Sınavı", "Türkçe", "İnkılap", "Din Kül.", "İngilizce", "Matematik", "Fen Bil."]
            for c_i, h in enumerate(sub_cmp_hdrs):
                c = tbl_sub_cmp.cell(0, c_i)
                c.paragraphs[0].text = h; c.paragraphs[0].runs[0].bold = True
                c.paragraphs[0].runs[0].font.size = Pt(7.5)
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_cell_shd(c, "1F497D"); c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
                set_cell_margins(c, 20, 20, 30, 30)

            r1_vals = ["1. PARAF LGS 1", f"{r1['tur_net']:.2f}", f"{r1['ink_net']:.2f}", f"{r1['din_net']:.2f}", f"{r1['ing_net']:.2f}", f"{r1['mat_net']:.2f}", f"{r1['fen_net']:.2f}"]
            r2_vals = ["2. NARTEST LGS 1", f"{r2['tur_net']:.2f}", f"{r2['ink_net']:.2f}", f"{r2['din_net']:.2f}", f"{r2['ing_net']:.2f}", f"{r2['mat_net']:.2f}", f"{r2['fen_net']:.2f}"]
            
            sub_keys = ["Türkçe", "İnkılap Tarihi", "Din Kültürü", "İngilizce", "Matematik", "Fen Bilimleri"]
            diff_vals = ["Net Farkı (+/-)"]
            for sk in sub_keys:
                df = st_comp['sub_diffs'].get(sk, 0.0)
                diff_vals.append(f"{df:+.2f}")

            for c_i, val in enumerate(r1_vals):
                cell = tbl_sub_cmp.cell(1, c_i)
                cell.paragraphs[0].text = val; cell.paragraphs[0].runs[0].font.size = Pt(7.5)
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_cell_margins(cell, 15, 15, 25, 25)

            for c_i, val in enumerate(diff_vals):
                cell = tbl_sub_cmp.cell(2, c_i)
                cell.paragraphs[0].text = val; cell.paragraphs[0].runs[0].font.size = Pt(7.5)
                cell.paragraphs[0].runs[0].bold = True
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_cell_margins(cell, 15, 15, 25, 25)
                if c_i > 0:
                    try:
                        f_val = float(val)
                        if f_val > 0: set_cell_shd(cell, "E8F8F5")
                        elif f_val < 0: set_cell_shd(cell, "FDEDEC")
                    except:
                        pass

            # Insight bullets
            p_ins = doc.add_paragraph()
            p_ins.paragraph_format.space_before = Pt(2)
            p_ins.paragraph_format.space_after = Pt(2)
            r_ins1 = p_ins.add_run(f"⭐ En Yüksek Artış: {st_comp['best_sub'][0]} ({st_comp['best_sub'][1]:+.2f} Net)  |  ")
            r_ins1.bold = True; r_ins1.font.size = Pt(8); r_ins1.font.color.rgb = RGBColor(14, 102, 85)
            r_ins2 = p_ins.add_run(f"⚠️ Takip / Müdahale Alanı: {st_comp['worst_sub'][0]} ({st_comp['worst_sub'][1]:+.2f} Net)")
            r_ins2.bold = True; r_ins2.font.size = Pt(8); r_ins2.font.color.rgb = RGBColor(180, 20, 20)

        elif st_comp and st_comp.get('status') == 'new_student':
            has_comp_section = True
            p_new = doc.add_paragraph()
            p_new.paragraph_format.space_before = Pt(4)
            p_new.paragraph_format.space_after = Pt(2)
            r_nh = p_new.add_run("🌟 1. YENİ KAYIT ÖĞRENCİ - REFERANS DENEME BİLGİLENDİRMESİ\n")
            r_nh.bold = True; r_nh.font.size = Pt(9); r_nh.font.color.rgb = RGBColor(31, 73, 125)
            r_nt = p_new.add_run(
                "Bu sınav (NARTEST LGS 1), öğrencimizin 8/B takip sistemimizdeki ilk referans (baz) denemesidir. "
                "Karşılaştırmalı net artış/azalış analizleri ve gelişim eğrisi bir sonraki deneme sınavından itibaren "
                "sistemimiz tarafından otomatik olarak sunulacaktır."
            )
            r_nt.font.size = Pt(8); r_nt.italic = True

    sec_num_k = "2" if has_comp_section else "1"
    sec_num_r = "3" if has_comp_section else "2"
    sec_num_a = "4" if has_comp_section else "3"

    # Diagnostic & Kazanım Summary
    diag = diagnose_student(student, kazanim_list)
    
    p_k = doc.add_paragraph()
    p_k.paragraph_format.space_before = Pt(3)
    p_k.paragraph_format.space_after = Pt(1)
    r = p_k.add_run(f"{sec_num_k}. GENEL KAZANIM BAŞARI VE EKSİK ENVANTERİ")
    r.bold = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(0, 102, 102)
    
    p_zero = doc.add_paragraph(style="List Bullet")
    p_zero.paragraph_format.space_after = Pt(1)
    r1 = p_zero.add_run("Öncelikli Tekrar Edilecek Kazanımlar (%0 Başarı): ")
    r1.bold = True; r1.font.size = Pt(8)
    txt_zero = ", ".join(diag['zero_kaz']) if diag['zero_kaz'] else "Tebrikler! Sıfır çeken kazanım bulunmuyor."
    r_z = p_zero.add_run(txt_zero)
    r_z.font.size = Pt(8)
    if diag['zero_kaz']: r_z.font.color.rgb = RGBColor(180, 20, 20)

    p_low = doc.add_paragraph(style="List Bullet")
    p_low.paragraph_format.space_after = Pt(1)
    r2 = p_low.add_run("Pekiştirilmesi Gereken Kritik Kazanımlar (<%50 Başarı): ")
    r2.bold = True; r2.font.size = Pt(8)
    txt_low = ", ".join(diag['low_kaz']) if diag['low_kaz'] else "Kritik eksik kazanım tespit edilmedi."
    r_l = p_low.add_run(txt_low)
    r_l.font.size = Pt(8)

    p_mast = doc.add_paragraph(style="List Bullet")
    p_mast.paragraph_format.space_after = Pt(2)
    r_m = p_mast.add_run("Tam Hakim Olunan Kazanımlar (%100 Başarı): ")
    r_m.bold = True; r_m.font.size = Pt(8)
    r_m_txt = p_mast.add_run(f"{len(diag['mastered_kaz'])} adet kazanımda %100 tam başarı sağlandı.")
    r_m_txt.font.size = Pt(8); r_m_txt.font.color.rgb = RGBColor(14, 102, 85)

    # 2. ALL 6 SUBJECTS DETAILED PRESCRIPTION (CRITICAL SECTION)
    prescriptions = get_subject_prescriptions(student, kazanim_list)

    p_rec_hdr = doc.add_paragraph()
    p_rec_hdr.paragraph_format.space_before = Pt(3)
    p_rec_hdr.paragraph_format.space_after = Pt(2)
    r_rh = p_rec_hdr.add_run(f"{sec_num_r}. TÜM DERSLER İÇİN KİŞİSELLEŞTİRİLMİŞ BRANŞ GELİŞİM REÇETELERİ")
    r_rh.bold = True; r_rh.font.size = Pt(9); r_rh.font.color.rgb = RGBColor(31, 73, 125)

    ordered_subs = ["Fen Bilimleri", "Matematik", "Türkçe", "İnkılap Tarihi", "Din Kültürü", "İngilizce"]

    for sub_name in ordered_subs:
        s_info = prescriptions.get(sub_name)
        if not s_info:
            continue
            
        tbl_card = doc.add_table(rows=2, cols=1)
        tbl_card.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl_card, color="B0C4DE", sz="4")
        
        # Header cell
        c_hdr = tbl_card.cell(0, 0)
        set_cell_shd(c_hdr, "F0F4F8")
        set_cell_margins(c_hdr, 30, 30, 50, 50)
        p_c_hdr = c_hdr.paragraphs[0]
        p_c_hdr.paragraph_format.space_after = Pt(0)
        r_sh = p_c_hdr.add_run(f"• {s_info['title'].upper()}  ")
        r_sh.bold = True; r_sh.font.size = Pt(8); r_sh.font.color.rgb = RGBColor(31, 73, 125)
        
        pct = (s_info['net']/20.0)*100 if sub_name in ["Türkçe", "Matematik", "Fen Bilimleri"] else (s_info['net']/10.0)*100
        r_sn = p_c_hdr.add_run(f"[Net: {s_info['net']:.2f} | D: {s_info['d']} | Y: {s_info['y']} | B: {s_info['b']} | Başarı: %{max(pct, 0.0):.1f}]")
        r_sn.bold = True; r_sn.font.size = Pt(7.5); r_sn.font.color.rgb = RGBColor(60, 60, 60)

        # Body cell
        c_body = tbl_card.cell(1, 0)
        set_cell_margins(c_body, 30, 30, 50, 50)
        
        # Line 1: Kazanım Dökümü
        p_b1 = c_body.paragraphs[0]
        p_b1.paragraph_format.space_after = Pt(1)
        r_k_lbl = p_b1.add_run("📊 Kazanım Başarı Dökümü: ")
        r_k_lbl.bold = True; r_k_lbl.font.size = Pt(7.5); r_k_lbl.font.color.rgb = RGBColor(0, 102, 102)
        r_k_txt = p_b1.add_run(s_info['summary'])
        r_k_txt.font.size = Pt(7.5)
        
        # Line 2: Sıfır Çekilenler (varsa)
        if s_info.get('zero'):
            p_b_zero = c_body.add_paragraph()
            p_b_zero.paragraph_format.space_after = Pt(1)
            r_z_l = p_b_zero.add_run("🔴 Sıfır Çekilen Kazanımlar (%0): ")
            r_z_l.bold = True; r_z_l.font.size = Pt(7.5); r_z_l.font.color.rgb = RGBColor(180, 20, 20)
            r_z_t = p_b_zero.add_run(", ".join(s_info['zero']))
            r_z_t.font.size = Pt(7.5)
            
        # Line 3: Pekiştirilecekler (varsa)
        if s_info.get('low'):
            p_b_low = c_body.add_paragraph()
            p_b_low.paragraph_format.space_after = Pt(1)
            r_lw_l = p_b_low.add_run("🟡 Pekiştirilecek Kazanımlar (<%100): ")
            r_lw_l.bold = True; r_lw_l.font.size = Pt(7.5); r_lw_l.font.color.rgb = RGBColor(180, 100, 0)
            r_lw_t = p_b_low.add_run(", ".join(s_info['low'][:5]))
            r_lw_t.font.size = Pt(7.5)
            
        # Line 4: Reçete
        p_b2 = c_body.add_paragraph()
        p_b2.paragraph_format.space_after = Pt(1)
        r_ad_lbl = p_b2.add_run("💡 Öğrenciye Özel Bireysel Çalışma Reçetesi: ")
        r_ad_lbl.bold = True; r_ad_lbl.font.size = Pt(7.5); r_ad_lbl.font.color.rgb = RGBColor(31, 73, 125)
        r_ad_txt = p_b2.add_run(s_info['advice'])
        r_ad_txt.font.size = Pt(7.5)

        # Line 5: Önerilen YouTube Desteği
        if s_info.get('youtube'):
            yt = s_info['youtube']
            p_b_yt = c_body.add_paragraph()
            p_b_yt.paragraph_format.space_after = Pt(1)
            r_yt_l = p_b_yt.add_run(f"📺 Önerilen YouTube Desteği ({yt['level_name']}): ")
            r_yt_l.bold = True; r_yt_l.font.size = Pt(7.5); r_yt_l.font.color.rgb = RGBColor(192, 57, 43)
            r_yt_c = p_b_yt.add_run(f"Kanal: {yt['channel']} | ")
            r_yt_c.bold = True; r_yt_c.font.size = Pt(7.5)
            r_yt_t = p_b_yt.add_run(f"Odak: {yt['focus']}. Taktik: {yt['study_tip']}")
            r_yt_t.font.size = Pt(7.5)
        
        # Tiny spacing between subject cards
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(0)
        p_sp.paragraph_format.space_after = Pt(1)

    # Action Targets
    p_act = doc.add_paragraph()
    p_act.paragraph_format.space_before = Pt(3)
    p_act.paragraph_format.space_after = Pt(1)
    r3 = p_act.add_run(f"{sec_num_a}. KİŞİSELLEŞTİRİLMİŞ HEDEFLER VE EYLEM TAKVİMİ")
    r3.bold = True; r3.font.size = Pt(9); r3.font.color.rgb = RGBColor(0, 102, 102)
    
    p_t14 = doc.add_paragraph(style="List Bullet")
    p_t14.paragraph_format.space_after = Pt(1)
    r_14 = p_t14.add_run("14. Gün Ara Kontrol Hedefi: "); r_14.bold = True; r_14.font.size = Pt(8)
    r_14_v = p_t14.add_run(f"{diag['target_14d']} Net (İlk hedefli konu tekrarları ve ödev kontrolü sonrası)")
    r_14_v.font.size = Pt(8)

    p_t28 = doc.add_paragraph(style="List Bullet")
    p_t28.paragraph_format.space_after = Pt(1)
    r_28 = p_t28.add_run("28. Gün (Sonraki Deneme) Hedefi: "); r_28.bold = True; r_28.font.size = Pt(8)
    r_28_v = p_t28.add_run(f"{diag['target_28d']} Net (Ölçülebilir bir sonraki resmi deneme sınavı hedefi)")
    r_28_v.font.size = Pt(8)

    p_rout = doc.add_paragraph(style="List Bullet")
    p_rout.paragraph_format.space_after = Pt(1)
    r_rt = p_rout.add_run("Haftalık Görev Kotası: "); r_rt.bold = True; r_rt.font.size = Pt(8)
    r_rt_v = p_rout.add_run(f"Haftalık {diag['weekly_quota']} soru çözümü + Yanlış Defteri (Hata Analizi) tutulması ve kontrolü.")
    r_rt_v.font.size = Pt(8)

    p_strat = doc.add_paragraph(style="List Bullet")
    p_strat.paragraph_format.space_after = Pt(3)
    r_st = p_strat.add_run("Öğrenciye Özel Genel Strateji: "); r_st.bold = True; r_st.font.size = Pt(8)
    r_st_v = p_strat.add_run(diag['strategy'])
    r_st_v.font.size = Pt(8)

    # Signatures
    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(4)
    tbl_sig = doc.add_table(rows=1, cols=3)
    tbl_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_sig, color="B0C4DE", sz="4")
    tbl_sig.cell(0, 0).paragraphs[0].text = "Öğrenci İmzası\n\n......................."
    tbl_sig.cell(0, 1).paragraphs[0].text = "Sınıf Rehber Öğretmeni\n\n......................."
    tbl_sig.cell(0, 2).paragraphs[0].text = "Veli Onayı ve İmzası\n\n......................."
    for c in range(3):
        tbl_sig.cell(0, c).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        tbl_sig.cell(0, c).paragraphs[0].runs[0].font.size = Pt(7.5)
        set_cell_margins(tbl_sig.cell(0, c), 35, 35, 45, 45)

    os.makedirs(os.path.dirname(os.path.abspath(output_file)), exist_ok=True)
    doc.save(output_file)

def generate_absent_student_card(student_info, exam_name, output_file, prev_exam_name="2026-2027  PARAF LGS 1"):
    doc = docx.Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.45); sec.bottom_margin = Inches(0.45)
    sec.left_margin = Inches(0.5); sec.right_margin = Inches(0.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"; normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(2); normal.paragraph_format.line_spacing = 1.15

    name = student_info['name'].strip()
    no = str(student_info.get('school_no') or student_info.get('num') or '---').strip()
    cls_name = student_info.get('class_name', '8/B')

    # Header
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run(f"{cls_name} SINIFI ÖĞRENCİ DENEME TAKİP RAPORU")
    r_t.bold = True; r_t.font.size = Pt(13); r_t.font.color.rgb = RGBColor(31, 73, 125)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_s = p_sub.add_run(f"Deneme Sınavı: {exam_name} | Sınav Katılım ve Telafi Takip Dosyası")
    r_s.italic = True; r_s.font.size = Pt(9); r_s.font.color.rgb = RGBColor(100, 100, 100)

    # Status Box
    tbl_st = doc.add_table(rows=2, cols=2)
    tbl_st.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_st, color="B0C4DE", sz="4")
    tbl_st.cell(0, 0).paragraphs[0].text = f"Öğrenci Adı Soyadı: {name} (No: {no} / {cls_name})"
    tbl_st.cell(0, 0).paragraphs[0].runs[0].bold = True
    tbl_st.cell(0, 1).paragraphs[0].text = "Sınav Katılım Durumu: KATILMADI (Mazeretli / Raporlu)"
    tbl_st.cell(0, 1).paragraphs[0].runs[0].bold = True
    tbl_st.cell(0, 1).paragraphs[0].runs[0].font.color.rgb = RGBColor(180, 20, 20)
    set_cell_shd(tbl_st.cell(0, 0), "F0F4F8")
    set_cell_shd(tbl_st.cell(0, 1), "FDEDEC")

    cell_desc = tbl_st.cell(1, 0)
    cell_desc.merge(tbl_st.cell(1, 1))
    p_d = cell_desc.paragraphs[0]
    r_d = p_d.add_run(
        f"Bilgilendirme: Öğrencimiz {name}, {exam_name} sınavına mazeretli/raporlu olarak katılamamıştır. "
        "Öğrencinin akademik takip dosyasında önceki resmi sınav başarıları eksiksiz korunmaktadır. "
        "Bu rapor öğrencinin sınavda sorulan kazanımlardan geri kalmaması ve konu telafisini tamamlaması amacıyla hazırlanmıştır."
    )
    r_d.font.size = Pt(8.5)
    set_cell_margins(cell_desc, 30, 30, 40, 40)

    # Previous exam baseline performance (from DB)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    p_prv = doc.add_paragraph()
    r_ph = p_prv.add_run("1. MEVCUT REFERANS AKADEMİK PERFORMANS (1. PARAF LGS 1)")
    r_ph.bold = True; r_ph.font.size = Pt(9.5); r_ph.font.color.rgb = RGBColor(31, 73, 125)

    from core.database import get_student_history
    hist = get_student_history(name)
    if hist:
        last_res = hist[0]
        tbl_p = doc.add_table(rows=2, cols=4)
        tbl_p.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl_p, color="B0C4DE", sz="4")
        p_labels = [
            ("Son Katılınan Sınav:", last_res['exam_name']),
            ("Kazanılan Toplam Net:", f"{last_res['top_net']:.2f} Net"),
            ("Elde Edilen LGS Puanı:", f"{last_res['puan']:.2f} Puan"),
            ("Sınıf / Okul Sıralaması:", f"{last_res['class_rank']}. Sıra / {last_res['school_rank']}. Sıra")
        ]
        for idx, (lbl, val) in enumerate(p_labels):
            r_i = idx // 2; c_i = (idx % 2) * 2
            c1 = tbl_p.cell(r_i, c_i); c2 = tbl_p.cell(r_i, c_i+1)
            c1.paragraphs[0].text = lbl; c1.paragraphs[0].runs[0].bold = True
            c1.paragraphs[0].runs[0].font.size = Pt(8)
            c2.paragraphs[0].text = val; c2.paragraphs[0].runs[0].font.size = Pt(8)
            set_cell_shd(c1, "DCE6F1")
            set_cell_margins(c1, 30, 30, 40, 40)
            set_cell_margins(c2, 30, 30, 40, 40)

    # Tested Topics Checklist
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    p_cat = doc.add_paragraph()
    r_ch = p_cat.add_run(f"2. {exam_name.upper()} SINAVINDA SORULAN VE TELAFİ EDİLECEK KONULAR LİSTESİ")
    r_ch.bold = True; r_ch.font.size = Pt(9.5); r_ch.font.color.rgb = RGBColor(0, 102, 102)

    tested_topics = [
        ("Türkçe (20 Soru)", "Fiilimsilerin İşlevi, Paragrafta Anlam ve Yapı, Ana Düşünce, Sözcükte/Cümlede Anlam, Deyim ve Atasözleri, Yazım Kuralları, Tablo ve Görsel Okuma, Söz Sanatları"),
        ("Matematik (20 Soru)", "Çarpanlar ve Katlar (Asal Çarpanlar), EBOB-EKOK Problemleri, Üslü İfadeler Temel Kuralları ve Negatif Üs"),
        ("Fen Bilimleri (20 Soru)", "Mevsimlerin Oluşumu (Eksen Eğikliği & Güneş Işınları), İklim ve Hava Hareketleri, DNA ve Genetik Kod, DNA Eşlenmesi"),
        ("T.C. İnkılap Tarihi (10 Soru)", "20. Yüzyıl Başlarında Osmanlı Devleti, Mustafa Kemal'in Çocukluk ve Öğrenim Hayatı, Fikir Hayatını Etkileyen Olaylar, Askerlik Hayatı"),
        ("Din Kültürü (10 Soru)", "Kader ve Kaza İnancı, İnsanın İradesi ve Kader, Kaderle İlgili Kavramlar (Ömür, Ecel, Rızık, Tevekkül)"),
        ("İngilizce (10 Soru)", "Friendship Ünitesi (Making Invitations, Accepting & Refusing, Making Excuses, Personal Qualities), The Internet")
    ]

    tbl_top = doc.add_table(rows=7, cols=2)
    tbl_top.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_top, color="B0C4DE", sz="4")
    tbl_top.cell(0, 0).paragraphs[0].text = "Branş Adı"; tbl_top.cell(0, 0).paragraphs[0].runs[0].bold = True
    tbl_top.cell(0, 1).paragraphs[0].text = "Sınavda Yer Alan Kazanım ve Konu Başlıkları"; tbl_top.cell(0, 1).paragraphs[0].runs[0].bold = True
    set_cell_shd(tbl_top.cell(0, 0), "1F497D"); set_cell_shd(tbl_top.cell(0, 1), "1F497D")
    tbl_top.cell(0, 0).paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    tbl_top.cell(0, 1).paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    set_cell_margins(tbl_top.cell(0, 0), 25, 25, 35, 35)
    set_cell_margins(tbl_top.cell(0, 1), 25, 25, 35, 35)

    for idx, (b_name, b_topics) in enumerate(tested_topics, 1):
        c_b = tbl_top.cell(idx, 0); c_t = tbl_top.cell(idx, 1)
        c_b.paragraphs[0].text = b_name; c_b.paragraphs[0].runs[0].bold = True
        c_b.paragraphs[0].runs[0].font.size = Pt(8)
        c_t.paragraphs[0].text = b_topics; c_t.paragraphs[0].runs[0].font.size = Pt(8)
        set_cell_margins(c_b, 20, 20, 30, 30)
        set_cell_margins(c_t, 20, 20, 30, 30)

    # Telafi Eylem Tavsiyesi
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    p_adv = doc.add_paragraph()
    r_ah = p_adv.add_run("3. AKADEMİK TELAFİ VE REHBERLİK ÖNERİLERİ\n")
    r_ah.bold = True; r_ah.font.size = Pt(9.5); r_ah.font.color.rgb = RGBColor(31, 73, 125)
    r_at = p_adv.add_run(
        "1. Öğrencimizin yukarıda listelenen kazanım konularında eksik kalmaması için MEB Kazanım Testlerini ve ilgili ünite sorularını evde çözmesi önerilir.\n"
        "2. Kaçırılan deneme sınavı kitapçığı okuldan temin edilerek süre tutularak (Sözel 75 dk, Sayısal 80 dk) ev ortamında çözülmeli ve öğretmenlere kontrol ettirilmelidir.\n"
        "3. Bir sonraki deneme sınavında akademik kondisyonu korumak için haftalık soru çözüm kotaları aksatılmamalıdır."
    )
    r_at.font.size = Pt(8.5)

    # Signatures
    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(4)
    tbl_sig = doc.add_table(rows=1, cols=3)
    tbl_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_sig, color="B0C4DE", sz="4")
    tbl_sig.cell(0, 0).paragraphs[0].text = "Öğrenci İmzası\n\n......................."
    tbl_sig.cell(0, 1).paragraphs[0].text = "Sınıf Rehber Öğretmeni\n\n......................."
    tbl_sig.cell(0, 2).paragraphs[0].text = "Veli Onayı ve İmzası\n\n......................."
    for c in range(3):
        tbl_sig.cell(0, c).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        tbl_sig.cell(0, c).paragraphs[0].runs[0].font.size = Pt(7.5)
        set_cell_margins(tbl_sig.cell(0, c), 35, 35, 45, 45)

    os.makedirs(os.path.dirname(os.path.abspath(output_file)), exist_ok=True)
    doc.save(output_file)

def generate_all_student_cards(students, exam_name, kazanim_data, output_dir, comp_data=None):
    os.makedirs(output_dir, exist_ok=True)
    created_files = []
    
    if comp_data is None:
        try:
            from core.comparative_analyzer import get_multi_exam_comparison
            comp_data = get_multi_exam_comparison()
        except Exception:
            comp_data = None
            
    # Clear old files in output_dir
    for existing in os.listdir(output_dir):
        if existing.endswith(".docx"):
            try:
                os.remove(os.path.join(output_dir, existing))
            except Exception:
                pass

    from core.database import get_all_students
    db_students = get_all_students()
    part_map = {s['name'].strip(): s for s in students}
    
    for dbs in db_students:
        name = dbs['name'].strip()
        no = str(dbs['school_no']).strip()
        safe_name = name.replace(" ", "_").replace("ç", "c").replace("Ç", "C").replace("ğ", "g").replace("Ğ", "G").replace("ı", "i").replace("İ", "I").replace("ö", "o").replace("Ö", "O").replace("ş", "s").replace("Ş", "S").replace("ü", "u").replace("Ü", "U")
        fpath = os.path.join(output_dir, f"{no}_{safe_name}_Karne.docx")
        
        if name in part_map:
            s_data = part_map[name]
            st_kaz = kazanim_data.get(name, []) if kazanim_data else []
            generate_single_student_card(s_data, exam_name, st_kaz, fpath, comp_data=comp_data)
        else:
            # Absent student
            generate_absent_student_card(dbs, exam_name, fpath)
            
        created_files.append(fpath)
        
    # Create ZIP package
    zip_path = os.path.join(os.path.dirname(output_dir), "8B_Tum_Karneler.zip")
    with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zf:
        for f in created_files:
            zf.write(f, os.path.basename(f))
            
    return created_files, zip_path
