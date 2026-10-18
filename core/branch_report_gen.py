import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os
import pandas as pd
import numpy as np
import zipfile

from config import SUBJECTS, KAZANIM_LEVELS, EXPORTS_DIR
from core.student_card_gen import clean_kaz_name, norm_text
from core.youtube_recommender import YOUTUBE_RESOURCES

# ===================== BRANŞ UZMANLIK VE PEDAGOJİ BİLGİ BANKASI =====================
BRANCH_PEDAGOGY = {
    "Matematik": {
        "title": "LGS Matematik Zümresi Akademik Değerlendirme ve Eylem Planı",
        "description": "20 Soru | Katsayı: 4.0 | LGS yerleştirme puanının ve Türkiye derecelerinin en belirleyici branşıdır.",
        "exam_characteristics": "LGS Matematik testi salt işlem bilgisi değil; okuduğunu anlama, matematiksel model kurma, veriyi tablo ve şekle dönüştürme ve zaman yönetimi becerisini ölçer.",
        "common_pitfalls": [
            "Yeni nesil sorularda uzun metin ve görseli görünce önyargıya kapılıp soruyu okumadan boş bırakma.",
            "EBOB ile EKOK sorularını ayırt edememe (Bütünden parçaya bölme/paylaştırma = EBOB; Parçadan bütüne katlama/birleşme = EKOK).",
            "Üslü ifadelerde parantez ve negatif üs yanılgıları (Örn: -2⁴ ile (-2)⁴ farkı; negatif üssün sayıyı ters çevirdiği kuralının unutulması).",
            "Aralarında asal sayılar ile asal sayı kavramını karıştırma (İki sayının aralarında asal olması için kendilerinin asal olması gerekmez).",
            "İşlem adımlarını kitapçık kenarına düzensiz yazıp işlem hatası yapma ve soru başına 2.5 dakikayı aşan inatlaşmalar."
        ],
        "tactics": [
            "3 Adımlı Problem Çözme: 1. Metindeki anahtar sayı ve birimleri yuvarlak içine al, 2. Verilenleri basit bir taslağa/şemaya dök, 3. İşlemi temiz yaz.",
            "Turlama Tekniği: 1. turda sadece hızlı anlaşılan ve doğrudan çözülen 8-12 soruyu çöz; zor ve karmaşık sorulara yıldız koyup 2. tura bırak.",
            "Negatif Net Engelleme: İki şık arasında kalınca tahmin yürütmek yerine boş bırak; 3 yanlışın 1 doğruyu götürdüğü bilincini koru."
        ],
        "four_week_plan": [
            ("1. Hafta (Kavram Onarımı & EBOB-EKOK):", "EBOB-EKOK ve Asal Çarpanlar konusunda sıfır çeken öğrencilere 15 dk mikro anlatım; akıllı tahtada şema ile modelleme; 10 temel + 10 yeni nesil soru çözümü."),
            ("2. Hafta (Üslü Sayılar Kurallar & İpuçları):", "Negatif üs, üslerin çarpımı/bölümü kural kartları dağıtımı; MEB çıkmış ve örnek sorularındaki çeldiricilerin adım adım analizi."),
            ("3. Hafta (Beceri Temelli Modelleme Atölyesi):", "Grup 1 ve 2 öğrencileri için ileri düzey yeni nesil problem çözümü; Grup 3 ve 4 öğrencileri için temel kavrama ve net artırma etüdü."),
            ("4. Hafta (Süre Baskılı Branş Provası):", "40 dakikalık 20 soruluk zaman sınırlı branş denemesi; Yanlış Defteri denetimi ve bireysel hedef revizyonu.")
        ],
        "weekly_homework_target": "Haftalık minimum 120-160 soru (Grup 1: 180+ soru; Grup 4: 90 temel/orta soru) + Hata Defteri analizi."
    },
    "Türkçe": {
        "title": "LGS Türkçe Zümresi Akademik Değerlendirme ve Eylem Planı",
        "description": "20 Soru | Katsayı: 4.0 | Sözel bölümün omurgasıdır; okuma hızı, odaklanma ve dil bilgisi mantığını ölçer.",
        "exam_characteristics": "LGS Türkçe testi paragrafta ana fikir, yardımcı düşünceler, sözel mantık, görsel/grafik yorumlama ve dil bilgisi (Fiilimsiler) kurallarını harmanlayan yüksek dikkat gerektiren bir sınavdır.",
        "common_pitfalls": [
            "Paragraf sorularında metnin tamamına odaklanmak yerine kendi ön bilgisini metne katarak çeldiriciye düşme.",
            "Olumsuz soru köklerini gözden kaçırma ('değinilmemiştir', 'çıkarılamaz', 'ulaşılamaz').",
            "Fiilimsileri çekimli fiillerle veya ek-fiillerle karıştırma; yüklemi fiilimsi sanma hatası.",
            "Öznel ve Nesnel anlatımlı cümlelerde kişisel kanı bildiren gizli sıfatları fark edememe.",
            "Sözel mantık sorularında şablon/tablo çizmeden akıldan eşleştirme yapıp süre kaybetme ve hata yapma."
        ],
        "tactics": [
            "Önce Soru Kökü: Paragrafı okumadan önce soru kökünü ve olumsuzluk ifadesini daire içine al.",
            "Kanıt Altı Çizme: Paragrafı okurken ana düşünceyi taşıyan ilk/son cümleleri ve anahtar sıfatları çiz.",
            "Fiilimsi Formülü: Cümlede önce yüklemi bul ve kapat; kalan sözcüklerde fiil kök/gövdesi arayıp -an,-ası,-mez / -mak,-me,-iş / -ken,-alı eklerini tespit et.",
            "Sözel Mantık Tablosu: Değişmeyen unsurları (günler, katlar, sıralar) sütun yap; değişkenleri (kişiler, nesneler) satırlara yerleştir."
        ],
        "four_week_plan": [
            ("1. Hafta (Fiilimsiler & Dil Bilgisi Tamamlama):", "Fiilimsi ekleri kodlama tablosu (Anası mezar dikecekmiş / Mayışmak / Kenyalı asiye...); 50 soruluk fiilimsi tarama testi."),
            ("2. Hafta (Cümlede Anlam & Öznel-Nesnel):", "Öznel-nesnel, sebep-sonuç, amaç-sonuç cümleleri ayrımı; MEB örnek sorularındaki gizli çeldiricilerin sınıfta tartışılması."),
            ("3. Hafta (Paragraf & Sözel Mantık Atölyesi):", "Grafik/tablo okuma ve sözel mantıkta şema kurma eğitimi; her dersin başında 2 adet yeni nesil sözel mantık çözümü."),
            ("4. Hafta (Zamanlı Paragraf & Hız Provası):", "20 soruluk Türkçe testini 25-30 dakikada bitirme denemesi; odaklanma ve soru kökü taktiklerinin pekiştirilmesi.")
        ],
        "weekly_homework_target": "Günlük 20-25 paragraf sorusu + haftalık 60 dil bilgisi/sözel mantık (Toplam: 180-200 soru)."
    },
    "Fen Bilimleri": {
        "title": "LGS Fen Bilimleri Zümresi Akademik Değerlendirme ve Eylem Planı",
        "description": "20 Soru | Katsayı: 4.0 | Sayısal bölümün net kaldıracıdır; deney, grafik ve kavram analizini ölçer.",
        "exam_characteristics": "LGS Fen Bilimleri testi bilimsel süreç becerilerini, kontrollü deney mantığını (bağımlı/bağımsız değişken), grafik ve tablo okumayı günlük hayat senaryoları üzerinden test eder.",
        "common_pitfalls": [
            "Deney sorularında bağımsız değişken (bizim değiştirdiğimiz), bağımlı değişken (etkilenen/ölçülen) ve sabit tutulan değişkeni karıştırma.",
            "Mevsimlerin oluşumunda Dünya'nın Güneş'e olan mesafesini sebep sanma yanılgısı (Temel neden: 23°27' eksen eğikliği ve Güneş ışınlarının geliş açısıdır).",
            "Güneş ışınlarının dik geldiği yerde gölge boyunun kısa, birim yüzeye düşen enerji miktarının fazla olduğunu unutma.",
            "Rüzgar oluşumunda yüksek basınç (soğuk alan) ile alçak basınç (sıcak alan) yönünü ters hatırlama (Rüzgar: Yüksek Basınç -> Alçak Basınç).",
            "Grafik eksenlerini (Sıcaklık, Zaman, Basınç, Gölge Boyu) dikkatle incelemeden hızlıca sonuca gitme."
        ],
        "tactics": [
            "Değişken Tablosu Çiz: Deney sorusunda 'Farklı olan ne?' -> Bağımsız Değişken. 'Ölçülen ne?' -> Bağımlı Değişken.",
            "Mevsim Pusulası: 21 Haziran (Yengeç dik, KYK Yaz), 21 Aralık (Oğlak dik, GYK Yaz), 21 Mart / 23 Eylül (Ekinoks, Ekvator dik).",
            "Basınç-Rüzgar Kodlaması: Soğuk Hava = Yoğun = Alçalır = Yüksek Basınç; Sıcak Hava = Seyrek = Yükselir = Alçak Basınç; Rüzgar = Soğuktan Sıcağa."
        ],
        "four_week_plan": [
            ("1. Hafta (Mevsimler ve Güneş Açıları Simülasyonu):", "3D simülasyon ve görsel animasyonlarla Dünya'nın konumu, eksen eğikliği ve gölge boyu ilişkisinin akıllı tahtada pekiştirilmesi."),
            ("2. Hafta (İklim, Hava Olayları & Basınç Deneyleri):", "Yüksek basınç - alçak basınç deney modelleri; rüzgar yönü ve nem/yağış türleri kavram haritası oluşturulması."),
            ("3. Hafta (DNA ve Genetik Kod Giriş & Deney Analizi):", "Bağımlı/bağımsız değişken soruları özel atölyesi; nükleotid dizilimi ve organik baz eşleşmeleri tarama testi."),
            ("4. Hafta (Grafik Yorumlama & 20 Soruluk Branş Denemesi):", "Karmaşık grafik ve çok aşamalı deney sorularından oluşan 40 dakikalık zaman sınırlı Fen denemesi.")
        ],
        "weekly_homework_target": "Haftalık 120-150 Fen sorusu (Özellikle deney ve grafik ağırlıklı yeni nesil sorular)."
    },
    "İnkılap Tarihi": {
        "title": "LGS T.C. İnkılap Tarihi ve Atatürkçülük Zümresi Değerlendirme ve Eylem Planı",
        "description": "10 Soru | Katsayı: 1.0 | Sözel bölümün net garantörüdür; kronoloji, harita ve kavram yorumunu ölçer.",
        "exam_characteristics": "LGS İnkılap Tarihi testi tarihsel metinleri yorumlama, neden-sonuç bağı kurma, harita ve kronoloji üzerinden Mustafa Kemal'in hayatı ve Milli Mücadele adımlarını analiz etmeyi gerektirir.",
        "common_pitfalls": [
            "'Yalnızca bu bilgiye göre' veya 'Metinden hareketle' sorularında kendi genel kültür bilgisini katarak metinde geçmeyen şıkkı işaretleme.",
            "Tarihsel kavramları birbirine karıştırma (Sömürgecilik = Hammadde ve pazar arayışı; Milliyetçilik = Kendi devletini kurma/isyan; Manda-Himaye = Başka devlete sığınma).",
            "Mustafa Kemal'in kişilik özellikleri ile olayları yanlış eşleştirme (Çanakkale: İleri görüşlülük/Askeri deha; Şam: Vatanseverlik/Liderlik; Trablusgarp: Teşkilatçılık).",
            "Avrupa'daki gelişmelerin (Sanayi İnkılabı, Fransız İhtilali) Osmanlı Devleti'ne siyasi ve ekonomik etkilerini ayırt edememe."
        ],
        "tactics": [
            "Metin Sınırında Kalma Kuralı: Soruda 'Yalnızca bu bilgiye göre' deniyorsa, şık doğru bir tarihi bilgi olsa bile metinde yazmıyorsa yanlıştır!",
            "Kavram-Eşleştirme Kartları: Sanayi İnkılabı -> Hammadde/Pazar -> Sömürgecilik -> Kapitülasyonlar -> Düyun-ı Umumiye zincirini ezberle.",
            "Kişilik Özellikleri Şifresi: Halkı örgütleme = Teşkilatçılık / Birleştiricilik; Önceden tahmin etme = İleri görüşlülük."
        ],
        "four_week_plan": [
            ("1. Hafta (1. Ünite Kavram Haritası & Kronoloji):", "Uyanan Avrupa, Sarsılan Osmanlı, Mustafa Kemal'in okuduğu okullar ve şehirler kronolojisinin harita üzerinde çizimi."),
            ("2. Hafta (Metin Yorumlama & 'Yalnızca' Kalıbı):", "'Metinden hareketle' soru kalıpları üzerinde sınıfta özel çalışma; çeldirici eleme stratejisi."),
            ("3. Hafta (Mustafa Kemal'in Askerlik Hayatı & Kişilik Özellikleri):", "Trablusgarp, Balkan Savaşları, Çanakkale Cephesi olay-kişilik özelliği eşleştirme testleri."),
            ("4. Hafta (10 Soruluk Hız Denemeleri & Full Hedefi):", "12-15 dakikalık mini denemelerle soru kaçırmama (10'da 10 yapma) kondisyonu kazanımı.")
        ],
        "weekly_homework_target": "Haftalık 60-80 İnkılap Tarihi sorusu + Kavram Defteri tekrarı."
    },
    "Din Kültürü": {
        "title": "LGS Din Kültürü ve Ahlak Bilgisi Zümresi Değerlendirme ve Eylem Planı",
        "description": "10 Soru | Katsayı: 1.0 | Sözel bölümde full yapılması hedeflenen; kavram ve ayet yorumuna dayalı derstir.",
        "exam_characteristics": "LGS Din Kültürü testi kaza-kader, evrensel yasalar, insanın iradesi, tevekkül ve ayet meallerini günlük hayat davranışlarıyla ilişkilendirmeyi ölçer.",
        "common_pitfalls": [
            "Evrensel yasaları (Fiziksel, Biyolojik, Toplumsal) birbiriyle karıştırma (Örn: Suyun kaldırma kuvveti = Fiziksel; Canlıların üremesi = Biyolojik; Adalet/Göç = Toplumsal).",
            "Tevekkül kavramını yanlış anlama (Gerekli tüm tedbirleri alıp çalıştıktan sonra Allah'a güvenmek yerine, hiç çalışmadan sonucu beklemek tevekkül değildir).",
            "Külli irade (Allah'ın sınırsız iradesi: doğum yeri, ırk, anne-baba) ile Cüzi irade (İnsanın seçimi: inanç, çalışkanlık, ahlak) sınırlarını ayırt edememe.",
            "Ayet el-Kürsi mealindeki Tevhid (Allah'ın birliği) ve kudret vurgularını gözden kaçırma.",
            "Ecel, ömür, rızık ve hastalık kavramlarını kaderci tembellikle bağdaştırma hatası."
        ],
        "tactics": [
            "Evrensel Yasa Şifresi: Cansız madde, enerji, yerçekimi, suyun kaynaması = Fiziksel; Canlı, bitki, hayvan, hücre, fotosentez = Biyolojik; İnsan ilişkileri, adalet, ahlak, eşitlik, göç = Toplumsal.",
            "Ayet-Kavram Eşleştirme: Ayetteki anahtar kelimeleri çiz (Gökleri ve yeri yaratan = Kudret/Fiziksel; Sizi bir damla sudan = Biyolojik; Haksızlık etmeyin = Toplumsal).",
            "10'da 10 Mantığı: Din Kültürü testinde 1 yanlış dahi LGS yüzdelik dilimini ciddi etkiler; şüpheli sorularda kavram tanımına geri dön."
        ],
        "four_week_plan": [
            ("1. Hafta (Kaza-Kader & Evrensel Yasalar Tasnifi):", "Fiziksel, biyolojik ve toplumsal yasaların 3 sütunlu kavram tablosunda 50 farklı örnek olayla sınıflandırılması."),
            ("2. Hafta (İnsanın İradesi & Tevekkül Analizi):", "Külli/Cüzi irade ayrımı; doğru tevekkül anlayışı üzerine örnek MEB soruları ve çeldirici analizi."),
            ("3. Hafta (Ayet el-Kürsi & Hz. Musa'nın Hayatı):", "Ayet el-Kürsi tefsir ve meal çalışması; Hz. Musa ve Firavun kıssasından çıkarılacak dersler tarama testi."),
            ("4. Hafta (10 Soruluk Branş Denemeleri & Sıfır Fire Provası):", "10 dakikalık 10 soruluk hız testleri; dikkatsizlik kaynaklı hataların sıfırlanması.")
        ],
        "weekly_homework_target": "Haftalık 50-70 soru + Ayet Mealleri okuma çalışması."
    },
    "İngilizce": {
        "title": "LGS İngilizce Zümresi Akademik Değerlendirme ve Eylem Planı",
        "description": "10 Soru | Katsayı: 1.0 | Sözel bölümde kelime bilgisi ve okuduğunu anlama üzerinden full hedeflenen branştır.",
        "exam_characteristics": "LGS İngilizce testi Türkçe çeviri değil; bağlamdan anlam çıkarma, diyalog akışını tamamlama, davet-ret-mazeret kalıpları ve görsel/davetiye kartlarını okumayı ölçer.",
        "common_pitfalls": [
            "Soru kökündeki olumsuzluk ifadelerini ve sınırlandırıcıları kaçırma ('Which of the following does NOT have an answer?', 'Who REFUSES the invitation?').",
            "Sadece reddeden kişi ile mazeret (excuse) bildiren kişiyi karıştırma ('I can't come' = Ret; 'Because I have an exam' = Mazeret).",
            "Friendship ünitesi sıfatlarını karıştırma (honest, reliable, generous, stubborn, jealous, supportive, aggressive).",
            "Davetiye kartlarındaki etkinlik türü (event), tarih (date), saat (time), yer (place/location), ücret (fee) ve gönderen/alıcı (sender/receiver) kavramlarını karıştırma."
        ],
        "tactics": [
            "Soru Kökünü Yuvarlak İçine Al: NOT, EXCEPT, REFUSE, ACCEPT, NO, TRUE, FALSE kelimelerini kesinlikle fosforlu kalemle belirle.",
            "Davetiye Analizi: Kim davet ediyor (sender - altta), kime yazılmış (receiver - üstte), ne zaman (date/time), nerede (place)?",
            "Mazeret Formülü: 'Because / I have to / My aunt is coming / I must study' ifadeleri mazeret (excuse) göstergesidir.",
            "Hedef 50 Kelime Kartı: Ünitedeki kritik sıfatları zıt anlamlılarıyla (honest != liar, generous != stingy) eşleştirip tekrar et."
        ],
        "four_week_plan": [
            ("1. Hafta (Friendship Ünitesi Kelime & Sıfat Avı):", "Hedef 50 kelime ve zıt anlamlı sıfatlar flashcard çalışması; kelime eşleştirme quiz uygulaması."),
            ("2. Hafta (Invitation, Accepting, Refusing & Excuses):", "Diyalog tamamlama teknikleri; teklif (offer), kabul (accept), ret (refuse) ve mazeret (excuse) ayırt etme atölyesi."),
            ("3. Hafta (Davetiye Kartı & Okuduğunu Anlama):", "Parti, konser, piknik davetiyeleri üzerinden soru-cevap çalışmaları; kim davet ediyor / kim geliyor analizi."),
            ("4. Hafta (10 Soruluk Hız Denemeleri & Soru Kökü Provası):", "10-12 dakikalık 10 soruluk İngilizce branş denemeleri; olumsuz soru köklerine sıfır hata disiplini.")
        ],
        "weekly_homework_target": "Haftalık 60-80 soru + Friendship Kelime Listesi günlük 10 dk sesli tekrar."
    }
}


# ===================== XML VE FORMATLAMA YARDIMCILARI =====================
def set_cell_shading(cell, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=70, bottom=70, left=100, right=100):
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>'
    )
    cell._tc.get_or_add_tcPr().append(tcMar)

def set_table_borders(table, color="D5D8DC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_callout(doc, text, alert_type="NOTE", title=None):
    colors = {
        "NOTE": ("1F497D", "F2F5F8"),
        "TIP": ("0E6655", "E8F8F5"),
        "IMPORTANT": ("7D6608", "FEF9E7"),
        "WARNING": ("900C3F", "FDEDEC")
    }
    border_col, bg_col = colors.get(alert_type, ("1F497D", "F2F5F8"))
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_shading(cell, bg_col)
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    borders_xml = (
        f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_col}"/>'
        f'<w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>'
    )
    cell._tc.get_or_add_tcPr().append(parse_xml(borders_xml))
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    
    tag_text = title if title else f"[{alert_type}]"
    r_tag = p.add_run(f"{tag_text}: ")
    r_tag.bold = True
    r_tag.font.color.rgb = RGBColor.from_string(border_col)
    r_tag.font.size = Pt(9.5)
    
    r_text = p.add_run(text)
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = RGBColor(50, 50, 50)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(1)
    p_after.paragraph_format.space_after = Pt(2)


# ===================== BRANŞ VERİ ANALİZİ =====================
def analyze_branch_data(subject_name, students, kazanim_data=None, comparisons=None):
    """
    Belirli bir branş için tüm sınıfın akademik, davranışsal ve kazanım analizini çıkarır.
    """
    # Eşleşme
    sub_key = subject_name
    for k in SUBJECTS.keys():
        if norm_text(k) in norm_text(subject_name) or norm_text(subject_name) in norm_text(k):
            sub_key = k
            break
            
    sub_info = SUBJECTS[sub_key]
    short = sub_info["short"]
    q_count = sub_info["questions"]

    d_col = f"{short}_d"
    y_col = f"{short}_y"
    b_col = f"{short}_b"
    net_col = f"{short}_net"

    student_list = []
    for s in students:
        s_d = int(s.get(d_col, 0) or 0)
        s_y = int(s.get(y_col, 0) or 0)
        s_b = int(s.get(b_col, 0) or 0)
        s_net = float(s.get(net_col, 0.0) or 0.0)
        top_net = float(s.get("top_net", 0.0) or 0.0)
        succ_pct = (s_net / q_count) * 100.0 if q_count > 0 else 0.0

        # Davranış analizi
        if s_y >= (q_count * 0.35) and s_y >= s_b * 1.3:
            behavior = "Yüksek Yanlış / Tahmin Eğilimi"
            behavior_code = "ATMA"
            recommendation = "İki şık arasında kalınca boş bırakma disiplini, kontrolsüz işaretlemeyi durdurma."
        elif s_b >= (q_count * 0.35):
            behavior = "Yüksek Boş / Süre ve Çekinme Eğilimi"
            behavior_code = "BOŞ"
            recommendation = "Soruya başlama cesareti, süre yönetimi ve temel kavram tekrarı."
        elif s_net >= (q_count * 0.80):
            behavior = "Zirve Hakimiyeti / Full Odaklı"
            behavior_code = "ZİRVE"
            recommendation = "MEB çeldirici analizleri, derece hedefi ve süreli zor soru çözümü."
        else:
            behavior = "Geliştirilebilir / Orta Seviye"
            behavior_code = "ORTA"
            recommendation = "Konu eksiğini kapatıp soru bankası taramasıyla neti yukarı taşıma."

        student_list.append({
            "name": s["name"],
            "school_no": s.get("school_no", s.get("num", "-")),
            "d": s_d,
            "y": s_y,
            "b": s_b,
            "net": s_net,
            "top_net": top_net,
            "score": float(s.get("score", s.get("puan", 0.0)) or 0.0),
            "succ_pct": succ_pct,
            "behavior": behavior,
            "behavior_code": behavior_code,
            "recommendation": recommendation
        })

    # Branş netine göre sırala
    student_list.sort(key=lambda x: (x["net"], x["d"], -x["y"]), reverse=True)
    for idx, st in enumerate(student_list, 1):
        st["branch_rank"] = idx

    df_b = pd.DataFrame(student_list)
    total_students = len(df_b)

    # İstatistikler
    mean_net = float(df_b["net"].mean()) if total_students > 0 else 0.0
    median_net = float(df_b["net"].median()) if total_students > 0 else 0.0
    max_net = float(df_b["net"].max()) if total_students > 0 else 0.0
    min_net = float(df_b["net"].min()) if total_students > 0 else 0.0
    std_net = float(df_b["net"].std()) if total_students > 0 else 0.0

    total_d = int(df_b["d"].sum()) if total_students > 0 else 0
    total_y = int(df_b["y"].sum()) if total_students > 0 else 0
    total_b = int(df_b["b"].sum()) if total_students > 0 else 0
    total_q_all = total_students * q_count

    d_pct = (total_d / total_q_all * 100.0) if total_q_all > 0 else 0.0
    y_pct = (total_y / total_q_all * 100.0) if total_q_all > 0 else 0.0
    b_pct = (total_b / total_q_all * 100.0) if total_q_all > 0 else 0.0

    # Okul Karşılaştırması
    school_avg = 0.0
    diff_from_school = 0.0
    if comparisons:
        for c in comparisons:
            if norm_text(c.get("subject", "")) in norm_text(sub_key) or norm_text(sub_key) in norm_text(c.get("subject", "")):
                school_avg = float(c.get("school_avg", c.get("school", 0.0)) or 0.0)
                diff_from_school = mean_net - school_avg
                break

    # Seviye Grupları (Branşa Özel Eşikler)
    is_major = (q_count == 20)
    t1_min = 16.0 if is_major else 8.0
    t2_min = 11.0 if is_major else 5.5
    t3_min = 6.0 if is_major else 3.0

    group_t1 = [s for s in student_list if s["net"] >= t1_min]
    group_t2 = [s for s in student_list if t2_min <= s["net"] < t1_min]
    group_t3 = [s for s in student_list if t3_min <= s["net"] < t2_min]
    group_t4 = [s for s in student_list if s["net"] < t3_min]

    tiers = {
        "Grup 1: Zirve & Full Hedefi (İleri Düzey)": {
            "range": f"{t1_min:.1f} - {q_count} Net",
            "students": group_t1,
            "goal": "Süre yönetimi, zorlayıcı MEB çeldirici analizi, Türkiye derecesi hazırlığı.",
            "color": "#1B4F72",
            "bg": "#EAECEE"
        },
        "Grup 2: Gelişim Liderleri (Potansiyel Grup)": {
            "range": f"{t2_min:.1f} - {t1_min - 0.1:.1f} Net",
            "students": group_t2,
            "goal": "Yeni nesil modelleme, işlem hatasını sıfırlama, düzenli soru bankası taraması.",
            "color": "#117864",
            "bg": "#E8F8F5"
        },
        "Grup 3: Kritik Müdahale (Destek Grubu)": {
            "range": f"{t3_min:.1f} - {t2_min - 0.1:.1f} Net",
            "students": group_t3,
            "goal": "Kazanım kavrama testleri, iki şık arasında kalınca boş bırakma taktiği.",
            "color": "#B7950B",
            "bg": "#FEF9E7"
        },
        "Grup 4: Temel Takviye (Bireysel İlgi)": {
            "range": f"0.0 - {t3_min - 0.1:.1f} Net",
            "students": group_t4,
            "goal": "Temel kavram konu videoları, kavram kartları ve negatif neti önleme.",
            "color": "#78281F",
            "bg": "#FDEDEC"
        }
    }

    # Kazanım Dökümü
    kazanim_summary = []
    if kazanim_data:
        all_k = []
        if isinstance(kazanim_data, dict):
            for sname, k_list in kazanim_data.items():
                for k in k_list:
                    sub_k = k.get("subject", "")
                    if norm_text(sub_k) in norm_text(sub_key) or norm_text(sub_key) in norm_text(sub_k):
                        c_name = clean_kaz_name(sub_key, k.get("kazanim") or k.get("kazanim_name", ""))
                        rate_v = float(k.get("rate") if k.get("rate") is not None else k.get("success_rate", 0.0))
                        all_k.append({"student_name": sname, "kazanim": c_name, "rate": rate_v})
        elif isinstance(kazanim_data, list):
            for k in kazanim_data:
                sub_k = k.get("subject", "")
                if norm_text(sub_k) in norm_text(sub_key) or norm_text(sub_key) in norm_text(sub_k):
                    c_name = clean_kaz_name(sub_key, k.get("kazanim") or k.get("kazanim_name", ""))
                    rate_v = float(k.get("rate") if k.get("rate") is not None else k.get("success_rate", 0.0))
                    all_k.append({"student_name": k.get("name") or k.get("student_name", ""), "kazanim": c_name, "rate": rate_v})

        if all_k:
            df_k = pd.DataFrame(all_k)
            for k_name, group in df_k.groupby("kazanim"):
                avg_rate = float(group["rate"].mean())
                zero_s = group[group["rate"] == 0.0]["student_name"].tolist()
                crit_s = group[group["rate"] < 40.0]["student_name"].tolist()
                full_s = group[group["rate"] == 100.0]["student_name"].tolist()
                high_s = group[group["rate"] >= 75.0]["student_name"].tolist()

                level_code = "D - GÜÇLÜ ALAN"
                level_desc = "Aralıklı tekrar ile formu koruma"
                level_color = "#0E6655"
                level_bg = "#E8F8F5"

                for kl in KAZANIM_LEVELS:
                    if avg_rate <= kl["max_rate"]:
                        level_code = kl["code"]
                        level_desc = kl["rec"]
                        level_color = kl["text_color"]
                        level_bg = kl["bg_color"]
                        break

                kazanim_summary.append({
                    "kazanim": k_name,
                    "avg_rate": avg_rate,
                    "level_code": level_code,
                    "level_desc": level_desc,
                    "level_color": level_color,
                    "level_bg": level_bg,
                    "zero_count": len(zero_s),
                    "zero_students": zero_s,
                    "crit_count": len(crit_s),
                    "crit_students": crit_s,
                    "full_count": len(full_s),
                    "full_students": full_s,
                    "high_count": len(high_s),
                    "high_students": high_s,
                    "total_count": len(group)
                })

            kazanim_summary.sort(key=lambda x: x["avg_rate"])

    pedagogy = BRANCH_PEDAGOGY.get(sub_key, BRANCH_PEDAGOGY["Matematik"])
    yt_resources = YOUTUBE_RESOURCES.get(sub_key, {})

    return {
        "subject": sub_key,
        "questions": q_count,
        "weight": sub_info["weight"],
        "total_students": total_students,
        "mean_net": mean_net,
        "median_net": median_net,
        "max_net": max_net,
        "min_net": min_net,
        "std_net": std_net,
        "school_avg": school_avg,
        "diff_from_school": diff_from_school,
        "total_d": total_d,
        "total_y": total_y,
        "total_b": total_b,
        "d_pct": d_pct,
        "y_pct": y_pct,
        "b_pct": b_pct,
        "students": student_list,
        "tiers": tiers,
        "kazanims": kazanim_summary,
        "pedagogy": pedagogy,
        "youtube": yt_resources
    }


# ===================== UZMAN BRANŞ WORD RAPORU OLUŞTURUCU =====================
def generate_branch_word_report(subject_name, exam_name, students, kazanim_data=None, comparisons=None, output_path=None):
    """
    Uzman bir branş öğretmeni bakış açısıyla MEB / Zümre teftiş standartlarında Word (.docx) raporu üretir.
    """
    b_data = analyze_branch_data(subject_name, students, kazanim_data, comparisons)
    ped = b_data["pedagogy"]
    real_subject = b_data["subject"]

    if not output_path:
        tr_safe = norm_text(real_subject).replace(" ", "_")
        output_path = os.path.join(EXPORTS_DIR, f"8B_{tr_safe.upper()}_Brans_Raporu.docx")

    doc = docx.Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(40, 40, 40)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.15

    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(15)
        r.font.color.rgb = RGBColor(31, 73, 125)

    def add_subtitle(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run(text)
        r.italic = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(110, 110, 110)

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(11.5)
        r.font.color.rgb = RGBColor(31, 73, 125)

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(17, 120, 100)

    # 1. BAŞLIK VE KÜNYE
    add_title(f"8/B SINIFI {real_subject.upper()} BRANŞI AKADEMİK ANALİZ VE EYLEM PLANI RAPORU")
    add_subtitle(f"Deneme Sınavı: {exam_name} | Branş: {real_subject} ({b_data['questions']} Soru - Katsayı: {b_data['weight']}x)\n"
                 f"Sınıf Mevcudu: {b_data['total_students']} Öğrenci | Analiz Motoru: LGS Otomasyonu v2.0")

    add_callout(
        doc,
        f"{ped['description']} {ped['exam_characteristics']}",
        "NOTE",
        title=f"UZMAN BRANŞ REHBERİ ({real_subject.upper()})"
    )

    # 2. BÖLÜM: BRANŞ GENEL PERFORMANS GÖSTERGELERİ
    add_h1("1. BRANŞ GENEL PERFORMANS GÖSTERGELERİ")

    tbl_kpi = doc.add_table(rows=2, cols=4)
    tbl_kpi.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_kpi)

    diff_str = f"+{b_data['diff_from_school']:.2f}" if b_data['diff_from_school'] >= 0 else f"{b_data['diff_from_school']:.2f}"
    kpi_items = [
        ("Sınıf Net Ortalaması", f"{b_data['mean_net']:.2f} / {b_data['questions']}", "F2F5F8"),
        ("Sınıf Medyan Neti", f"{b_data['median_net']:.2f}", "F2F5F8"),
        ("Okul Ortalaması / Fark", f"{b_data['school_avg']:.2f} ({diff_str})", "E8F8F5" if b_data['diff_from_school']>=0 else "FDEDEC"),
        ("Zirve / Taban Net", f"{b_data['max_net']:.2f} / {b_data['min_net']:.2f}", "F2F5F8"),
        ("Toplam Doğru", f"{b_data['total_d']} (%{b_data['d_pct']:.1f})", "E8F8F5"),
        ("Toplam Yanlış", f"{b_data['total_y']} (%{b_data['y_pct']:.1f})", "FDEDEC"),
        ("Toplam Boş", f"{b_data['total_b']} (%{b_data['b_pct']:.1f})", "FEF9E7"),
        ("Standart Sapma", f"{b_data['std_net']:.2f} Net", "F2F5F8"),
    ]

    for idx, (label, val, bg) in enumerate(kpi_items):
        r_idx = idx // 4
        c_idx = idx % 4
        cell = tbl_kpi.cell(r_idx, c_idx)
        set_cell_shading(cell, bg)
        set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p.add_run(f"{label}\n")
        r1.font.size = Pt(8)
        r1.font.color.rgb = RGBColor(100, 100, 100)
        r2 = p.add_run(val)
        r2.bold = True
        r2.font.size = Pt(11)
        r2.font.color.rgb = RGBColor(31, 73, 125)

    doc.add_paragraph().paragraph_format.space_before = Pt(4)

    # Branş Yanlış / Boş Davranış Tespiti
    if b_data['total_y'] > b_data['total_b']:
        add_callout(
            doc,
            f"8/B sınıfında {real_subject} testinde toplam YANLIŞ sayısı ({b_data['total_y']}), BOŞ sayısından ({b_data['total_b']}) fazladır. "
            f"Öğrencilerde emin olmadan işaretleme (şans/tahmin) davranışı gözlenmektedir. LGS'de 3 yanlış 1 doğruyu götürdüğü için "
            f"bu durum sınıfın net ortalamasını yaklaşık {(b_data['total_y']/3)/b_data['total_students']:.2f} net aşağı çekmektedir. "
            f"Öğrencilere 'İki güçlü çeldirici arasında kalındığında boş bırakma' disiplini verilmelidir.",
            "WARNING",
            title="SINAV STRATEJİSİ UYARISI: NEGATİF NET KAYBI"
        )
    else:
        add_callout(
            doc,
            f"8/B sınıfında {real_subject} testinde toplam BOŞ sayısı ({b_data['total_b']}), YANLIŞ sayısından ({b_data['total_y']}) fazladır. "
            f"Öğrenciler sorularda süre yetiştirememiş veya yeni nesil uzun metinlerden çekinmiş olabilir. "
            f"Zaman yönetimi ve soruya cesaretle başlama pratikleri yapılmalıdır.",
            "TIP",
            title="SINAV STRATEJİSİ TAVSİYESİ: SÜRE VE CESARET YÖNETİMİ"
        )

    # 3. BÖLÜM: BRANŞA ÖZEL EN SIK YAPILAN KAVRAM HATALARI
    add_h1(f"2. {real_subject.upper()} TESTİNDE EN SIK DÜŞÜLEN TUZAKLAR VE KAVRAM YANILGILARI")
    for pit in ped["common_pitfalls"]:
        p_pit = doc.add_paragraph(style="List Bullet")
        p_pit.paragraph_format.space_before = Pt(1)
        p_pit.paragraph_format.space_after = Pt(1)
        r_w = p_pit.add_run("⚠️ ")
        r_t = p_pit.add_run(pit)
        r_t.font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_before = Pt(3)
    add_h2(f"💡 Uzman {real_subject} Öğretmeni Çözüm Taktikleri")
    for tac in ped["tactics"]:
        p_tac = doc.add_paragraph(style="List Bullet")
        p_tac.paragraph_format.space_before = Pt(1)
        p_tac.paragraph_format.space_after = Pt(1)
        r_w = p_tac.add_run("✔ ")
        r_w.bold = True
        r_w.font.color.rgb = RGBColor(14, 102, 85)
        r_t = p_tac.add_run(tac)
        r_t.font.size = Pt(9)

    # 4. BÖLÜM: KAZANIM BAZLI BAŞARI VE ACİL MÜDAHALE LİSTESİ
    doc.add_paragraph().paragraph_format.space_before = Pt(4)
    add_h1(f"3. {real_subject.upper()} KAZANIM BAŞARI DÖKÜMÜ VE ÖNCELİKLENDİRME")
    doc.add_paragraph(
        f"Aşağıdaki tabloda {real_subject} dersine ait kazanımların sınıf ortalamaları, başarı düzeyleri ve "
        f"zümre öğretmenine önerilen ders içi müdahale reçetesi yer almaktadır:"
    )

    if b_data["kazanims"]:
        tbl_kaz = doc.add_table(rows=len(b_data["kazanims"]) + 1, cols=6)
        tbl_kaz.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl_kaz)

        k_hdrs = ["Kazanım Adı", "Başarı %", "Öncelik Kodu", "Kritik Açık (<%40)", "Güçlü Hakim (≥%75)", "Zümre Reçetesi"]
        for c, h in enumerate(k_hdrs):
            cell = tbl_kaz.cell(0, c)
            set_cell_shading(cell, "1F497D")
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            p.text = h
            p.runs[0].bold = True
            p.runs[0].font.size = Pt(8.5)
            p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        for r_idx, k_item in enumerate(b_data["kazanims"], 1):
            row_cells = tbl_kaz.rows[r_idx].cells
            
            bg_col = "FFFFFF"
            if k_item["avg_rate"] < 40.0:
                bg_col = "FDEDEC"
            elif k_item["avg_rate"] < 60.0:
                bg_col = "FEF9E7"
            elif k_item["avg_rate"] >= 80.0:
                bg_col = "E8F8F5"

            vals = [
                k_item["kazanim"],
                f"%{k_item['avg_rate']:.1f}",
                k_item["level_code"].split(" - ")[0],
                f"{k_item['crit_count']} Öğr.",
                f"{k_item['high_count']} Öğr.",
                k_item["level_desc"]
            ]

            for c_idx, val in enumerate(vals):
                cell = row_cells[c_idx]
                set_cell_shading(cell, bg_col)
                set_cell_margins(cell, top=45, bottom=45, left=65, right=65)
                p = cell.paragraphs[0]
                p.text = str(val)
                p.runs[0].font.size = Pt(8)
                if c_idx in [1, 2, 3, 4]:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.runs[0].bold = True

        # Kritik desteğe ihtiyacı olan öğrencilerin net isim listesi
        doc.add_paragraph().paragraph_format.space_before = Pt(3)
        add_h2("📌 Kritik Kazanımlarda Desteğe İhtiyacı Olan Öğrenciler (<%40 Başarı / Sıfır Çekenler)")
        crit_found = False
        for k_item in b_data["kazanims"]:
            if k_item["crit_count"] > 0 and k_item["avg_rate"] < 65.0:
                crit_found = True
                p_z = doc.add_paragraph(style="List Bullet")
                p_z.paragraph_format.space_before = Pt(1)
                p_z.paragraph_format.space_after = Pt(1)
                r_z1 = p_z.add_run(f"{k_item['kazanim']} (Sınıf Başarısı: %{k_item['avg_rate']:.1f}): ")
                r_z1.bold = True
                r_z1.font.size = Pt(8.5)
                r_z1.font.color.rgb = RGBColor(144, 12, 63)
                r_z2 = p_z.add_run(f"Kritik {k_item['crit_count']} Öğrenci: {', '.join(k_item['crit_students'])}")
                r_z2.font.size = Pt(8.5)
        if not crit_found:
            doc.add_paragraph("Tebrikler! Kritik eşikte (<%40) desteğe muhtaç öğrenci grubu bulunmamaktadır.")
    else:
        doc.add_paragraph("Bu branş için kayıtlı kazanım dökümü bulunmamaktadır.")

    # 5. BÖLÜM: FARKLILAŞTIRILMIŞ ÖĞRETİM İÇİN ÖĞRENCİ SEVİYE GRUPLARI
    doc.add_paragraph().paragraph_format.space_before = Pt(4)
    add_h1(f"4. FARKLILAŞTIRILMIŞ ÖĞRETİM GRUPLARI ({real_subject.upper()} BAZLI)")
    doc.add_paragraph(
        f"Öğrenciler {real_subject} netlerine göre 4 gruba tasnif edilmiştir. "
        f"Zümre öğretmeninin ders içi ödevlendirme ve etüt planında bu grupları esas alması önerilir:"
    )

    for g_title, g_info in b_data["tiers"].items():
        tbl_g = doc.add_table(rows=1, cols=1)
        tbl_g.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell_g = tbl_g.cell(0, 0)
        set_cell_shading(cell_g, g_info["bg"][1:] if g_info["bg"].startswith("#") else g_info["bg"])
        set_cell_margins(cell_g, top=70, bottom=70, left=110, right=110)
        border_hex = g_info["color"][1:] if g_info["color"].startswith("#") else g_info["color"]
        borders_xml = (
            f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/>'
            f'<w:left w:val="single" w:sz="20" w:space="0" w:color="{border_hex}"/>'
            f'<w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>'
        )
        cell_g._tc.get_or_add_tcPr().append(parse_xml(borders_xml))
        
        p = cell_g.paragraphs[0]
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        r_head = p.add_run(f"{g_title} [{g_info['range']}] — {len(g_info['students'])} Öğrenci\n")
        r_head.bold = True
        r_head.font.size = Pt(9.5)
        r_head.font.color.rgb = RGBColor.from_string(border_hex)

        r_goal = p.add_run(f"🎯 Zümre Hedefi & Strateji: {g_info['goal']}\n")
        r_goal.italic = True
        r_goal.font.size = Pt(8.5)

        st_names_str = ", ".join([f"{s['name']} ({s['net']:.2f}N)" for s in g_info["students"]]) if g_info["students"] else "Bu grupta öğrenci bulunmuyor."
        r_st = p.add_run(f"Öğrenciler: {st_names_str}")
        r_st.font.size = Pt(8.5)

        doc.add_paragraph().paragraph_format.space_before = Pt(2)

    # 6. BÖLÜM: TÜM ÖĞRENCİLERİN BRANŞ PERFORMANS VE DAVRANIŞ SIRALAMASI
    add_h1(f"5. 8/B SINIFI {real_subject.upper()} ÖĞRENCİ SIRALAMA VE DAVRANIŞ LİSTESİ")

    tbl_st = doc.add_table(rows=len(b_data["students"]) + 1, cols=8)
    tbl_st.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_st)

    s_hdrs = ["Sıra", "No", "Öğrenci Adı Soyadı", "D", "Y", "B", "Net", "Davranış / Bireysel Tavsiye"]
    for c, h in enumerate(s_hdrs):
        cell = tbl_st.cell(0, c)
        set_cell_shading(cell, "595959")
        set_cell_margins(cell, top=45, bottom=45, left=60, right=60)
        p = cell.paragraphs[0]
        p.text = h
        p.runs[0].bold = True
        p.runs[0].font.size = Pt(8)
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for r_idx, s in enumerate(b_data["students"], 1):
        row_cells = tbl_st.rows[r_idx].cells
        
        bg_col = "FFFFFF" if r_idx % 2 == 1 else "F9F9F9"
        if s["net"] >= (b_data["questions"] * 0.8):
            bg_col = "E8F8F5"
        elif s["net"] < (b_data["questions"] * 0.3):
            bg_col = "FDEDEC"

        vals = [
            str(s["branch_rank"]),
            str(s["school_no"]),
            s["name"],
            str(s["d"]),
            str(s["y"]),
            str(s["b"]),
            f"{s['net']:.2f}",
            f"{s['behavior_code']}: {s['recommendation']}"
        ]

        for c_idx, val in enumerate(vals):
            cell = row_cells[c_idx]
            set_cell_shading(cell, bg_col)
            set_cell_margins(cell, top=35, bottom=35, left=50, right=50)
            p = cell.paragraphs[0]
            p.text = str(val)
            p.runs[0].font.size = Pt(7.5)
            if c_idx in [0, 1, 3, 4, 5, 6]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if c_idx == 6:
                p.runs[0].bold = True

    # 7. BÖLÜM: 4 HAFTALIK ZÜMRE EYLEM VE İYİLEŞTİRME TAKVİMİ
    doc.add_paragraph().paragraph_format.space_before = Pt(6)
    add_h1(f"6. {real_subject.upper()} ZÜMRESİ 4 HAFTALIK EYLEM VE TAKİP PLANI")
    doc.add_paragraph(f"Önümüzdeki deneme sınavına kadar derste uygulanacak somut eylem adımları:")

    for week, act in ped["four_week_plan"]:
        p_act = doc.add_paragraph(style="List Bullet")
        p_act.paragraph_format.space_before = Pt(1)
        p_act.paragraph_format.space_after = Pt(1)
        r_w = p_act.add_run(f"{week} ")
        r_w.bold = True
        r_w.font.size = Pt(8.5)
        r_w.font.color.rgb = RGBColor(31, 73, 125)
        r_a = p_act.add_run(act)
        r_a.font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_before = Pt(2)
    p_hw = doc.add_paragraph()
    r_hw1 = p_hw.add_run("🎯 Haftalık Ödevlendirme ve Soru Çözüm Kotası: ")
    r_hw1.bold = True
    r_hw1.font.size = Pt(9)
    r_hw1.font.color.rgb = RGBColor(17, 120, 100)
    r_hw2 = p_hw.add_run(ped["weekly_homework_target"])
    r_hw2.font.size = Pt(9)

    # 8. BÖLÜM: SEVİYE BAZLI YOUTUBE DESTEK REHBERİ
    if b_data.get("youtube"):
        doc.add_paragraph().paragraph_format.space_before = Pt(2)
        add_h2(f"📺 {real_subject} Öğrencilere Tavsiye Edilecek YouTube Kanalları")
        for lvl_key, l_info in b_data["youtube"].items():
            p_yt = doc.add_paragraph(style="List Bullet")
            p_yt.paragraph_format.space_before = Pt(1)
            p_yt.paragraph_format.space_after = Pt(1)
            r_yt1 = p_yt.add_run(f"{l_info['level_name']}: {l_info['channel']} — ")
            r_yt1.bold = True
            r_yt1.font.size = Pt(8.5)
            r_yt2 = p_yt.add_run(f"Odak: {l_info['focus']} | Taktiği: {l_info['study_tip']}")
            r_yt2.font.size = Pt(8.5)

    # 9. BÖLÜM: RESMİ İMZA VE ONAY ALANI
    doc.add_paragraph().paragraph_format.space_before = Pt(10)
    tbl_sign = doc.add_table(rows=2, cols=3)
    tbl_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    sign_headers = [f"{real_subject} Öğretmeni", "Zümre Başkanı", "Okul Müdürü"]
    for i, h in enumerate(sign_headers):
        c0 = tbl_sign.cell(0, i)
        c0.paragraphs[0].text = h
        c0.paragraphs[0].runs[0].bold = True
        c0.paragraphs[0].runs[0].font.size = Pt(9)
        c0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        c1 = tbl_sign.cell(1, i)
        c1.paragraphs[0].text = "\n\nİmza / Tarih: ................"
        c1.paragraphs[0].runs[0].font.size = Pt(8)
        c1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    if output_path:
        out_dir = os.path.dirname(os.path.abspath(output_path))
        os.makedirs(out_dir, exist_ok=True)

    doc.save(output_path)
    return output_path


def generate_all_branch_reports(exam_name, students, kazanim_data=None, comparisons=None, output_dir=None):
    """
    Tüm 6 LGS branşı için uzman Word raporu üretir ve tek bir ZIP paketi haline getirir.
    """
    if not output_dir:
        output_dir = os.path.join(EXPORTS_DIR, "brans_raporlari")
    os.makedirs(output_dir, exist_ok=True)

    generated_files = []
    for sub in SUBJECTS.keys():
        tr_safe = norm_text(sub).replace(" ", "_")
        f_name = f"8B_{tr_safe.upper()}_Brans_Raporu.docx"
        out_p = os.path.join(output_dir, f_name)
        generate_branch_word_report(sub, exam_name, students, kazanim_data, comparisons, out_p)
        generated_files.append(out_p)

    # ZIP Paketi
    zip_path = os.path.join(EXPORTS_DIR, "8B_Tum_Brans_Ogretmen_Raporlari.zip")
    with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zf:
        for f in generated_files:
            zf.write(f, os.path.basename(f))

    return generated_files, zip_path
