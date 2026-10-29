import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os
import sys

root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if root not in sys.path:
    sys.path.insert(0, root)

if sys.stdout.encoding and sys.stdout.encoding.lower() != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

from config import DATABASE_PATH, EXPORTS_DIR
from core.database import get_exam_results, get_student_kazanims, get_class_comparison, list_exams, get_all_students
from core.analyzer import analyze_class_performance, calculate_pareto, diagnose_student, generate_support_groups
from core.comparative_analyzer import get_multi_exam_comparison

def build_master_excel(output_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active) # Remove default sheet

    # Color Palette
    navy_header = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    navy_light = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    accent_teal = PatternFill(start_color="008080", end_color="008080", fill_type="solid")
    gray_header = PatternFill(start_color="595959", end_color="595959", fill_type="solid")
    zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    alert_red = PatternFill(start_color="FADBD8", end_color="FADBD8", fill_type="solid")
    warn_yellow = PatternFill(start_color="FCF3CF", end_color="FCF3CF", fill_type="solid")
    good_green = PatternFill(start_color="D4EFDF", end_color="D4EFDF", fill_type="solid")
    mastered_green = PatternFill(start_color="D5F5E3", end_color="D5F5E3", fill_type="solid")

    font_title = Font(name="Calibri", size=15, bold=True, color="1F497D")
    font_sub = Font(name="Calibri", size=10, italic=True, color="595959")
    font_h1 = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_reg = Font(name="Calibri", size=9.5)

    thin_border = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                         top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))

    # Load Exams
    exams = list_exams()
    paraf_exam = exams[0]
    nartest_exam = exams[1]

    # Load Active (NARTEST) Data
    nartest_students = get_exam_results(nartest_exam['id'])
    sorted_students = sorted(nartest_students, key=lambda x: x.get('class_rank', 999))
    kazanim_data = {s['name']: get_student_kazanims(nartest_exam['id'], s['name']) for s in nartest_students}
    comparisons = get_class_comparison(nartest_exam['id'])
    comp_map = {c.get('subject'): c.get('school_avg', 0.0) for c in comparisons}

    stats = analyze_class_performance(nartest_students)
    pareto = calculate_pareto(nartest_students)
    groups = generate_support_groups(nartest_students, kazanim_data)

    # Load Multi-Exam Comparisons
    multi_comp = get_multi_exam_comparison()

    # =========================================================================
    # SHEET 1: Denemeler_Arası_Gelişim (MASTER COMPARATIVE SHEET)
    # =========================================================================
    ws_comp = wb.create_sheet(title="Denemeler_Arası_Gelişim")
    ws_comp.views.sheetView[0].showGridLines = True
    ws_comp.cell(1, 1, "8/B SINIFI DENEMELER ARASI AKADEMİK GELİŞİM VE NET DEĞİŞİM TAKİP TABLOSU").font = font_title
    ws_comp.cell(2, 1, f"1. Deneme: {paraf_exam['name']}  ➔  2. Deneme: {nartest_exam['name']} | Toplam Sınıf Mevcudu: 29 Öğrenci").font = font_sub

    # KPI Sınıf Gelişim Özeti
    ws_comp.cell(4, 1, "SINIF DÜZEYİNDE DENEMELER ARASI DEĞİŞİM VE GENEL BAŞARI").font = Font(name="Calibri", size=11, bold=True, color="008080")
    kpi_cmp_hdrs = ["Metrik Adı", "1. PARAF LGS 1 (12.10)", "2. NARTEST LGS 1 (26.10)", "Sınavlar Arası Fark", "Değerlendirme ve Trend"]
    for c, h in enumerate(kpi_cmp_hdrs, 1):
        cell = ws_comp.cell(5, c, h); cell.font = font_h1; cell.fill = accent_teal; cell.alignment = Alignment(horizontal="center")

    kpi_cmp_rows = [
        ("Sınıf Net Ortalaması", "37.43 Net", "44.80 Net", "+7.37 Net Artış", "🚀 Çok Güçlü İlerleme"),
        ("Sınıf Puan Ortalaması", "304.00 Puan", "335.31 Puan", "+31.31 Puan Artış", "📈 Belirgin Puan Artışı"),
        ("En Yüksek Sınıf Neti", "80.68 Net (Orhan Efe)", "83.67 Net (Orhan Efe)", "+2.99 Net Artış", "Zirve Seviyesi Yükseldi"),
        ("Sınıf Medyan Neti", "37.00 Net", "44.16 Net", "+7.16 Net Artış", "Sınıf Tabanı Güçlendi"),
        ("Okul Genel Net Ortalaması", "36.50 Net", "45.80 Net", "+9.30 Net Artış", "Okul Ortalaması Takibi"),
    ]
    for r_idx, rdata in enumerate(kpi_cmp_rows, 6):
        for c_idx, val in enumerate(rdata, 1):
            cell = ws_comp.cell(r_idx, c_idx, val); cell.font = font_reg; cell.border = thin_border
            cell.alignment = Alignment(horizontal="left" if c_idx == 1 else "center")
            if c_idx == 4:
                cell.fill = good_green; cell.font = font_bold

    # Öğrenci Bazlı Karşılaştırmalı Liste
    ws_comp.cell(12, 1, "ÖĞRENCİ BAZINDA NET VE PUAN DEĞİŞİM LİSTESİ (GELİŞİM SIRASINA GÖRE)").font = Font(name="Calibri", size=11, bold=True, color="1F497D")
    stu_cmp_hdrs = [
        "Sıra", "No", "Öğrenci Adı Soyadı", "Giriş PIN",
        "Paraf Net", "Nartest Net", "Net Farkı",
        "Paraf Puan", "Nartest Puan", "Puan Farkı",
        "Sınıf Sıra", "Okul Sıra",
        "ΔTürkçe", "Δİnkılap", "ΔDin", "Δİngilizce", "ΔMatematik", "ΔFen",
        "En Çok Gelişen Branş", "Takip Gereken Branş", "Genel Durum"
    ]
    for c, h in enumerate(stu_cmp_hdrs, 1):
        cell = ws_comp.cell(13, c, h); cell.font = font_h1; cell.fill = navy_header; cell.alignment = Alignment(horizontal="center")

    # Sort students: valid first by delta_net desc, then new/absent
    comp_list = list(multi_comp.values())
    valid_cmp = [c for c in comp_list if c['status'] == 'normal']
    valid_cmp.sort(key=lambda x: x['delta_net'], reverse=True)
    other_cmp = [c for c in comp_list if c['status'] != 'normal']
    full_cmp_list = valid_cmp + other_cmp

    for r_idx, c in enumerate(full_cmp_list, 14):
        st_no = c['school_no']
        st_name = c['name']
        pin = c['pin']
        
        if c['status'] == 'normal':
            r1 = c['r1']; r2 = c['r2']
            d_net = c['delta_net']
            d_puan = c['delta_puan']
            d_cls = c['delta_cls_rank']
            d_sch = c['delta_sch_rank']
            diff_cls_str = f"{d_cls:+d}" if d_cls != 0 else "0"
            diff_sch_str = f"{d_sch:+d}" if d_sch != 0 else "0"
            
            s_diffs = c['sub_diffs']
            best_str = f"{c['best_sub'][0]} ({c['best_sub'][1]:+.2f})"
            worst_str = f"{c['worst_sub'][0]} ({c['worst_sub'][1]:+.2f})"
            status_badge = f"{c['trend_icon']} {c['trend']}"
            
            vals = [
                r_idx - 13, st_no, st_name, pin,
                r1['top_net'], r2['top_net'], d_net,
                r1['puan'], r2['puan'], d_puan,
                f"{r1['class_rank']}➔{r2['class_rank']} ({diff_cls_str})",
                f"{r1['school_rank']}➔{r2['school_rank']} ({diff_sch_str})",
                s_diffs.get('Türkçe', 0.0), s_diffs.get('İnkılap Tarihi', 0.0), s_diffs.get('Din Kültürü', 0.0),
                s_diffs.get('İngilizce', 0.0), s_diffs.get('Matematik', 0.0), s_diffs.get('Fen Bilimleri', 0.0),
                best_str, worst_str, status_badge
            ]
        elif c['status'] == 'new_student':
            r2 = c['r2']
            vals = [
                r_idx - 13, st_no, st_name, pin,
                "-", r2['top_net'], "-",
                "-", r2['puan'], "-",
                f"-➔{r2['class_rank']}", f"-➔{r2['school_rank']}",
                "-", "-", "-", "-", "-", "-",
                "-", "-", "🌟 Yeni Kayıt (Baz Deneme)"
            ]
        else: # absent
            r1 = c['r1']
            vals = [
                r_idx - 13, st_no, st_name, pin,
                r1['top_net'] if r1 else "-", "-", "-",
                r1['puan'] if r1 else "-", "-", "-",
                f"{r1['class_rank']}➔-" if r1 else "-", f"{r1['school_rank']}➔-" if r1 else "-",
                "-", "-", "-", "-", "-", "-",
                "-", "-", "💤 Katılmadı (Mazeretli)"
            ]

        for c_idx, v in enumerate(vals, 1):
            cell = ws_comp.cell(r_idx, c_idx, v); cell.font = font_reg; cell.border = thin_border
            cell.alignment = Alignment(horizontal="left" if c_idx in [3, 19, 20, 21] else "center")
            if c_idx == 3: cell.font = font_bold
            
            # Net Farkı Color Coding
            if c_idx == 7 and isinstance(v, (int, float)):
                if v > 0: cell.fill = good_green; cell.font = font_bold
                elif v < 0: cell.fill = alert_red; cell.font = font_bold
            # Puan Farkı Color Coding
            elif c_idx == 10 and isinstance(v, (int, float)):
                if v > 0: cell.fill = good_green; cell.font = font_bold
                elif v < 0: cell.fill = alert_red; cell.font = font_bold
            # Branch differences
            elif c_idx in range(13, 19) and isinstance(v, (int, float)):
                if v > 0: cell.fill = good_green
                elif v < 0: cell.fill = alert_red

    # =========================================================================
    # SHEET 2: Sınıf_Genel_Görünümü (NARTEST 1 DASHBOARD)
    # =========================================================================
    ws1 = wb.create_sheet(title="Sınıf_Genel_Görünümü")
    ws1.views.sheetView[0].showGridLines = True
    ws1.cell(1, 1, f"8/B SINIFI LGS AKADEMİK GELİŞİM VE TAKİP PANELİ - {nartest_exam['name']}").font = font_title
    ws1.cell(2, 1, f"Toplam Katılım: {stats.get('count', 0)} Öğrenci | Okul Genelinde 56 Katılımcı | Kurumsal Başarı Takip Sistemi").font = font_sub

    # Table 1: KPI Overview
    kpi_headers = ["Toplam Öğrenci", "Net Ortalaması", "Medyan Net", "En Yüksek Net", "En Düşük Net", "Puan Ortalaması", "Standart Sapma"]
    for c, h in enumerate(kpi_headers, 1):
        cell = ws1.cell(4, c, h); cell.font = font_h1; cell.fill = navy_header; cell.alignment = Alignment(horizontal="center")
    kpi_vals = [
        stats.get('count', len(nartest_students)),
        round(stats.get('mean_net', 0.0), 2),
        round(stats.get('median_net', 0.0), 2),
        round(stats.get('max_net', 0.0), 2),
        round(stats.get('min_net', 0.0), 2),
        round(stats.get('mean_score', 0.0), 2),
        round(stats.get('std_net', 0.0), 2)
    ]
    for c, v in enumerate(kpi_vals, 1):
        cell = ws1.cell(5, c, v); cell.font = font_bold; cell.border = thin_border; cell.alignment = Alignment(horizontal="center")
        cell.fill = navy_light

    # Table 2: Subject Performance & School Comparison
    ws1.cell(7, 1, "DERS BAZLI SINIF PERFORMANSI VE OKUL ORTALAMALARI KARŞILAŞTIRMASI").font = Font(name="Calibri", size=11, bold=True, color="008080")
    sub_hdrs = ["Ders Adı", "Soru Sayısı", "Doğru Ort.", "Yanlış Ort.", "Boş Ort.", "8/B Net Ort.", "Okul Net Ort.", "Fark (+/-)", "Durum / Yorum"]
    for c, h in enumerate(sub_hdrs, 1):
        cell = ws1.cell(8, c, h); cell.font = font_h1; cell.fill = accent_teal; cell.alignment = Alignment(horizontal="center")

    n_students = max(len(nartest_students), 1)
    sub_rows = [
        ("Türkçe", 20, sum(s.get('tur_d',0) for s in nartest_students)/n_students, sum(s.get('tur_y',0) for s in nartest_students)/n_students, sum(s.get('tur_b',0) for s in nartest_students)/n_students, sum(s.get('tur_net',0.0) for s in nartest_students)/n_students),
        ("İnkılap Tarihi", 10, sum(s.get('ink_d',0) for s in nartest_students)/n_students, sum(s.get('ink_y',0) for s in nartest_students)/n_students, sum(s.get('ink_b',0) for s in nartest_students)/n_students, sum(s.get('ink_net',0.0) for s in nartest_students)/n_students),
        ("Din Kültürü", 10, sum(s.get('din_d',0) for s in nartest_students)/n_students, sum(s.get('din_y',0) for s in nartest_students)/n_students, sum(s.get('din_b',0) for s in nartest_students)/n_students, sum(s.get('din_net',0.0) for s in nartest_students)/n_students),
        ("İngilizce", 10, sum(s.get('ing_d',0) for s in nartest_students)/n_students, sum(s.get('ing_y',0) for s in nartest_students)/n_students, sum(s.get('ing_b',0) for s in nartest_students)/n_students, sum(s.get('ing_net',0.0) for s in nartest_students)/n_students),
        ("Matematik", 20, sum(s.get('mat_d',0) for s in nartest_students)/n_students, sum(s.get('mat_y',0) for s in nartest_students)/n_students, sum(s.get('mat_b',0) for s in nartest_students)/n_students, sum(s.get('mat_net',0.0) for s in nartest_students)/n_students),
        ("Fen Bilimleri", 20, sum(s.get('fen_d',0) for s in nartest_students)/n_students, sum(s.get('fen_y',0) for s in nartest_students)/n_students, sum(s.get('fen_b',0) for s in nartest_students)/n_students, sum(s.get('fen_net',0.0) for s in nartest_students)/n_students),
        ("TOPLAM NET", 90, sum(s.get('top_d',0) for s in nartest_students)/n_students, sum(s.get('top_y',0) for s in nartest_students)/n_students, sum(s.get('top_b',0) for s in nartest_students)/n_students, sum(s.get('top_net',0.0) for s in nartest_students)/n_students)
    ]

    for r_idx, sdata in enumerate(sub_rows, 9):
        s_name, q_cnt, avg_d, avg_y, avg_b, avg_net = sdata
        sch_avg = comp_map.get(s_name, 0.0)
        diff = avg_net - sch_avg if sch_avg > 0 else 0.0
        status = "Okul Üstü (+)" if diff > 0 else ("Okul Altı (-)" if diff < 0 else "Okul ile Eşit")
        row_vals = [s_name, q_cnt, round(avg_d, 2), round(avg_y, 2), round(avg_b, 2), round(avg_net, 2), round(sch_avg, 2) if sch_avg else "-", f"{diff:+.2f}" if sch_avg else "-", status]
        for c_idx, v in enumerate(row_vals, 1):
            cell = ws1.cell(r_idx, c_idx, v); cell.font = font_reg; cell.border = thin_border
            cell.alignment = Alignment(horizontal="left" if c_idx in [1, 9] else "center")
            if r_idx == 15:
                cell.font = font_bold; cell.fill = navy_light
            elif c_idx == 8 and sch_avg > 0:
                if diff > 0: cell.fill = good_green
                elif diff < 0: cell.fill = alert_red

    # Table 3: Pareto Net Loss Table
    ws1.cell(17, 1, "PARETO ANALİZİ - NET KAYBININ BRANŞLARA DAĞILIMI VE ÖNCELİKLİ MÜDAHALE").font = Font(name="Calibri", size=11, bold=True, color="1F497D")
    pareto_hdrs = ["Öncelik Sırası", "Ders Adı", "Kayıp Net (Sınıf Toplamı)", "Toplam Yanlış", "Toplam Boş", "Kayıptaki Payı (%)", "Kümülatif Pay (%)", "Müdahale Düzeyi"]
    for c, h in enumerate(pareto_hdrs, 1):
        cell = ws1.cell(18, c, h); cell.font = font_h1; cell.fill = navy_header; cell.alignment = Alignment(horizontal="center")

    for r_idx, p in enumerate(pareto, 19):
        lvl = "Kritik Acil Müdahale" if p['cum_pct'] <= 50 else ("Öncelikli Gelişim" if p['cum_pct'] <= 80 else "İzleme ve Koruma")
        p_vals = [p['rank'], p['subject'], round(p['lost_net'], 2), p['total_y'], p['total_b'], f"%{p['share_pct']:.2f}", f"%{p['cum_pct']:.2f}", lvl]
        for c_idx, v in enumerate(p_vals, 1):
            cell = ws1.cell(r_idx, c_idx, v); cell.font = font_reg; cell.border = thin_border
            cell.alignment = Alignment(horizontal="left" if c_idx in [2, 8] else "center")
            if c_idx == 8:
                if "Kritik" in lvl: cell.fill = alert_red; cell.font = font_bold
                elif "Öncelikli" in lvl: cell.fill = warn_yellow
                else: cell.fill = good_green

    # =========================================================================
    # SHEET 3: Öğrenci_Sonuçları_ve_Sıralama (NARTEST RESULTS)
    # =========================================================================
    ws2 = wb.create_sheet(title="Öğrenci_Sonuçları_ve_Sıralama")
    ws2.views.sheetView[0].showGridLines = True
    ws2.cell(1, 1, f"8/B SINIFI ÖĞRENCİ SONUÇ VE SIRALAMA LİSTESİ ({nartest_exam['name']})").font = font_title
    headers_s2 = ["S.Sıra", "No", "Öğrenci Adı Soyadı", "T_Net", "İ_Net", "D_Net", "E_Net", "M_Net", "F_Net", "Top_D", "Top_Y", "Top_B", "Top_Net", "Puanı", "Okul Sıra", "TG Sıra", "TG % Dilim", "Ceza Kaybı", "Kademe"]
    for c, h in enumerate(headers_s2, 1):
        cell = ws2.cell(3, c, h); cell.font = font_h1; cell.fill = navy_header; cell.alignment = Alignment(horizontal="center")

    for r_idx, s in enumerate(sorted_students, 4):
        t_net = s.get('top_net', 0.0)
        tier = "Tier 1 (Hedef: Fen Lisesi)" if t_net >= 65 else ("Tier 2 (Hedef: Anadolu Lisesi)" if t_net >= 40 else ("Tier 3 (Hedef: Temel Takviye)" if t_net >= 25 else "Tier 4 (Hedef: Acil Destek)"))
        pen = round(s.get('top_y', 0) / 3.0, 2)
        vals = [
            s.get('class_rank', r_idx-3), s.get('school_no', '-'), s.get('name'),
            s.get('tur_net', 0.0), s.get('ink_net', 0.0), s.get('din_net', 0.0), s.get('ing_net', 0.0), s.get('mat_net', 0.0), s.get('fen_net', 0.0),
            s.get('top_d', 0), s.get('top_y', 0), s.get('top_b', 0), t_net, s.get('puan', 0.0),
            s.get('school_rank', '-'), s.get('tg_rank', '-'), f"%{s.get('dilim'):.2f}" if s.get('dilim') else "-", pen, tier
        ]
        for c_idx, v in enumerate(vals, 1):
            cell = ws2.cell(r_idx, c_idx, v); cell.font = font_reg; cell.border = thin_border
            cell.alignment = Alignment(horizontal="left" if c_idx in [3, 19] else "center")
            if c_idx == 3: cell.font = font_bold
            if c_idx == 13: cell.font = font_bold; cell.fill = navy_light

    # =========================================================================
    # SHEET 4: Öğrenci_Kazanım_Matrisi
    # =========================================================================
    if kazanim_data:
        ws3 = wb.create_sheet(title="Öğrenci_Kazanım_Matrisi")
        ws3.views.sheetView[0].showGridLines = True
        ws3.cell(1, 1, f"8/B SINIFI ÖĞRENCİ - KAZANIM BAŞARI MATRİSİ ({nartest_exam['name']})").font = font_title
        
        first_st = list(kazanim_data.keys())[0]
        kaz_defs = [(k['subject'], k['kazanim']) for k in kazanim_data[first_st]]
        
        ws3.cell(3, 1, "Ders").font = font_h1; ws3.cell(3, 1).fill = navy_header
        ws3.cell(3, 2, "Kazanım Adı").font = font_h1; ws3.cell(3, 2).fill = navy_header
        
        for idx, s in enumerate(sorted_students):
            col = 3 + idx
            cell = ws3.cell(3, col, f"{s.get('class_rank')}. {s['name'][:12]}")
            cell.font = font_h1; cell.fill = navy_header; cell.alignment = Alignment(horizontal="center", wrap_text=True)

        for r_idx, (sub, kaz) in enumerate(kaz_defs, 4):
            ws3.cell(r_idx, 1, sub).border = thin_border
            ws3.cell(r_idx, 2, kaz).border = thin_border
            for idx, s in enumerate(sorted_students):
                col = 3 + idx
                name = s['name']
                st_k_list = kazanim_data.get(name, [])
                rate = 0.0
                for item in st_k_list:
                    if item['subject'] == sub and item['kazanim'] == kaz:
                        rate = item['rate']
                        break
                cell = ws3.cell(r_idx, col, f"%{rate:.0f}")
                cell.font = font_reg; cell.border = thin_border; cell.alignment = Alignment(horizontal="center")
                if rate == 100.0: cell.fill = mastered_green
                elif rate >= 50.0: cell.fill = warn_yellow
                else: cell.fill = alert_red

    # =========================================================================
    # SHEET 5: Bireysel_Hedefler_ve_Takip
    # =========================================================================
    ws4 = wb.create_sheet(title="Bireysel_Hedefler_ve_Takip")
    ws4.views.sheetView[0].showGridLines = True
    ws4.cell(1, 1, "8/B SINIFI ÖĞRENCİ HEDEF KARTLARI VE KAZANIM EKSİK REHBERİ").font = font_title
    headers_s4 = ["Sıra", "No", "Öğrenci Adı Soyadı", "Mevcut Net", "Puan", "TG Sıra", "14. Gün Hedef", "28. Gün Hedef", "Sıfır Çeken Kazanım", "Haftalık Soru Kotası", "Stratejik Odak"]
    for c, h in enumerate(headers_s4, 1):
        cell = ws4.cell(3, c, h); cell.font = font_h1; cell.fill = accent_teal; cell.alignment = Alignment(horizontal="center")

    for r_idx, s in enumerate(sorted_students, 4):
        diag = diagnose_student(s, kazanim_data.get(s['name'], []) if kazanim_data else None)
        vals = [
            s.get('class_rank', r_idx-3), s.get('school_no', '-'), s.get('name'), s.get('top_net', 0.0), s.get('puan', 0.0), s.get('tg_rank', '-'),
            diag['target_14d'], diag['target_28d'], len(diag['zero_kaz']), diag['weekly_quota'], diag['strategy']
        ]
        for c_idx, v in enumerate(vals, 1):
            cell = ws4.cell(r_idx, c_idx, v); cell.font = font_reg; cell.border = thin_border
            cell.alignment = Alignment(horizontal="left" if c_idx in [3, 11] else "center")
            if c_idx == 3: cell.font = font_bold
            if c_idx in [7, 8]: cell.fill = good_green; cell.font = font_bold

    # =========================================================================
    # SHEET 6: Destek_Grupları
    # =========================================================================
    ws5 = wb.create_sheet(title="Destek_Grupları")
    ws5.views.sheetView[0].showGridLines = True
    ws5.cell(1, 1, "8/B SINIFI İHTİYAÇ ODAKLI DİNAMİK DESTEK GRUPLARI").font = font_title
    headers_s5 = ["Grup No ve Adı", "Öğrenci Sayısı", "Dahil Edilen Öğrenciler", "Haftalık Süre"]
    for c, h in enumerate(headers_s5, 1):
        cell = ws5.cell(3, c, h); cell.font = font_h1; cell.fill = navy_header; cell.alignment = Alignment(horizontal="center")
    for r_idx, (gname, members) in enumerate(groups.items(), 4):
        vals = [gname, len(members), ", ".join(members) if members else "Öğrenci yok", "2-3 Saat/Hafta"]
        for c_idx, v in enumerate(vals, 1):
            cell = ws5.cell(r_idx, c_idx, v); cell.font = font_reg; cell.border = thin_border
            cell.alignment = Alignment(horizontal="left" if c_idx in [1, 3] else "center")
            if c_idx == 1: cell.font = font_bold; cell.fill = navy_light

    # Auto adjust column widths across all sheets
    for sheet in wb.worksheets:
        for col in sheet.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                val = cell.value
                if val:
                    for l in str(val).split('\n'):
                        if len(l) > max_len: max_len = len(l)
            sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 45)

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)
    return output_path

if __name__ == "__main__":
    out_desktop = r"C:\Users\ELRA\Desktop\8B sınıfı veriler\8B_LGS_Takip_Sistemi.xlsx"
    out_exports = os.path.join(EXPORTS_DIR, "8B_LGS_Takip_Sistemi.xlsx")
    
    print(f"Building Master Excel Tracker...")
    build_master_excel(out_desktop)
    build_master_excel(out_exports)
    print(f"Saved Master Excel to: {out_desktop} ({os.path.getsize(out_desktop):,} bytes)")
    print(f"Saved Master Excel to: {out_exports} ({os.path.getsize(out_exports):,} bytes)")
