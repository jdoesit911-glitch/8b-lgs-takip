import os
import docx
from docx.shared import Inches, Pt

reports_dir = r"C:\Users\ELRA\Desktop\8B sınıfı veriler\brans_raporlari"
files = [
    "8B_TURKCE_Brans_Raporu.docx",
    "8B_MATEMATIK_Brans_Raporu.docx",
    "8B_FEN_BILIMLERI_Brans_Raporu.docx",
    "8B_INKILAP_TARIHI_Brans_Raporu.docx",
    "8B_DIN_KULTURU_Brans_Raporu.docx",
    "8B_INGILIZCE_Brans_Raporu.docx"
]

print("=== VERIFYING 6 BRANCH TEACHER REPORTS ===")
for f in files:
    p = os.path.join(reports_dir, f)
    d = docx.Document(p)
    sec = d.sections[0]
    w_in = round(sec.page_width.inches, 2)
    h_in = round(sec.page_height.inches, 2)
    orient = "Landscape" if sec.orientation == 1 or w_in > h_in else "Portrait"
    
    m_top = round(sec.top_margin.inches, 3)
    m_bot = round(sec.bottom_margin.inches, 3)
    m_l = round(sec.left_margin.inches, 3)
    m_r = round(sec.right_margin.inches, 3)
    
    t_kaz = d.tables[1]
    widths = []
    # check grid or cell widths
    grid = t_kaz._tbl.tblPr.xpath('.//w:tblGrid/w:gridCol/@w:w')
    
    # check 4th col font size in table 1
    font_sizes = []
    for row in t_kaz.rows[1:]:
        c = row.cells[3]
        for run in c.paragraphs[0].runs:
            if run.font.size:
                font_sizes.append(run.font.size.pt)
                
    fs_sample = set(font_sizes)
    
    print(f"\nFile: {f}")
    print(f" Size: {w_in} x {h_in} in ({orient}) | Margins: T={m_top}, B={m_bot}, L={m_l}, R={m_r}")
    print(f" Kazanım Table Col Widths (dxa): {grid}")
    print(f" 4th Col Font Sizes (Pt): {fs_sample}")
